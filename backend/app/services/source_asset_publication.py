"""Bounded, mechanical evidence that exported source images are public.

Pinning bytes locally and serving those bytes at the configured public origin
are separate facts. Only this application's configured HTTPS origin and exact
content-addressed source-asset route are eligible for a probe. Model-provided
external URLs, redirects and arbitrary query parameters are never fetched.
Reports are frozen at staging; workbook rendering performs no network I/O.
"""
from __future__ import annotations

import hashlib
import io
import ipaddress
import json
import os
import re
import time
from datetime import datetime, timezone
from typing import Any, Mapping
from urllib.parse import urlsplit

import httpx

from .. import config
from . import source_asset_store

REPORT_FIELD = "source_asset_publication"
VERSION = "source-asset-publication-1"
MAX_ASSET_BYTES = 8 * 1024 * 1024
MAX_IMAGE_PIXELS = 25_000_000
MAX_PROBES = 128
TOTAL_SECONDS = 20.0
REQUEST_SECONDS = 2.0

# These parse the public wire grammar; they never decide visual ownership.
_IMAGE = re.compile(r'\[img\s+src="([^"\n]+)"\s+alt="[^"\n]*"\]')
_SOURCE_URL = re.compile(r'https://[^\s<>"\'\[\]()]+/source-assets/[^\s<>"\'\[\]()]+')
_ROUTE = re.compile(r"/source-assets/([0-9]+)/([0-9a-f]{64}\.jpg)")


def configured_origin() -> str:
    return str(os.environ.get("AEGIS_PUBLIC_BASE_URL", config.PUBLIC_BASE_URL) or "").strip().rstrip("/")


def valid_public_origin(value: str) -> bool:
    """Validate an operator-supplied origin, never trust an origin from content."""
    try:
        parsed = urlsplit(value)
        host = str(parsed.hostname or "").lower()
        if (
            parsed.scheme != "https" or not host or parsed.path not in {"", "/"}
            or parsed.query or parsed.fragment or parsed.username or parsed.password
            or parsed.port not in {None, 443}
            or host == "localhost" or host.endswith((".localhost", ".local", ".invalid"))
        ):
            return False
        try:
            return ipaddress.ip_address(host).is_global
        except ValueError:
            return "." in host
    except ValueError:
        return False


def _origin(value: str) -> tuple[str, str, int | None]:
    parsed = urlsplit(value)
    return parsed.scheme.lower(), str(parsed.hostname or "").lower(), parsed.port or 443


def image_urls(value: Any) -> list[str]:
    """Collect image wires, not arbitrary links or model-selected network targets."""
    found: set[str] = set()

    def walk(item: Any) -> None:
        if isinstance(item, Mapping):
            # An Image-typed cell contains one bare image URL, not rich markup.
            medium = str(item.get("answer_type") or item.get("type") or "").lower()
            if medium == "image":
                for field in ("answer_content", "content", "keyword"):
                    url = str(item.get(field) or "").strip()
                    if url.startswith("https://"):
                        found.add(url)
            for key, child in item.items():
                if not str(key).startswith("_"):
                    walk(child)
        elif isinstance(item, (list, tuple)):
            for child in item:
                walk(child)
        elif isinstance(item, str):
            found.update(_IMAGE.findall(item))
            found.update(_SOURCE_URL.findall(item))

    walk(value)
    return sorted(found)


def _input_sha(urls: list[str]) -> str:
    return hashlib.sha256(json.dumps(urls, separators=(",", ":")).encode()).hexdigest()


def release_asset_input(payload: Mapping[str, Any]) -> dict[str, Any]:
    """Only exported content; source inventory and audit quotes are not outputs."""
    return {
        "records": payload.get("records") or [],
        "chapter_meta": payload.get("chapter_meta") or {},
    }


def _local_status(filename: str) -> str:
    try:
        path = source_asset_store.stored_asset_path(filename)
        if not path.is_file():
            return "not_pinned"
        if path.stat().st_size > MAX_ASSET_BYTES:
            return "size_not_verified"
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        return "pinned" if digest == filename.removesuffix(".jpg") else "hash_mismatch"
    except OSError:
        return "unreadable"


def _probe(client: httpx.Client, url: str, expected_hash: str, deadline: float) -> dict[str, Any]:
    try:
        with client.stream("GET", url, follow_redirects=False) as response:
            if response.status_code != 200:
                return {"state": "failed", "reason": "http_status", "http_status": response.status_code}
            if response.headers.get("content-type", "").split(";", 1)[0].strip().lower() != "image/jpeg":
                return {"state": "failed", "reason": "content_type"}
            declared = response.headers.get("content-length")
            if declared and declared.isdigit() and int(declared) > MAX_ASSET_BYTES:
                return {"state": "failed", "reason": "response_too_large"}
            digest = hashlib.sha256()
            size = 0
            chunks: list[bytes] = []
            # Do not accumulate a fixed-size chunk before checking the total
            # deadline: a slow trickle must not evade an inactivity timeout.
            for chunk in response.iter_bytes():
                if time.monotonic() >= deadline:
                    return {"state": "unverified", "reason": "probe_budget_exhausted"}
                size += len(chunk)
                if size > MAX_ASSET_BYTES:
                    return {"state": "failed", "reason": "response_too_large"}
                digest.update(chunk)
                chunks.append(chunk)
            if digest.hexdigest() != expected_hash:
                return {"state": "failed", "reason": "content_hash_mismatch", "bytes": size}
            from PIL import Image

            try:
                with Image.open(io.BytesIO(b"".join(chunks))) as decoded:
                    if decoded.format != "JPEG":
                        return {"state": "failed", "reason": "not_a_jpeg_image"}
                    if decoded.width * decoded.height > MAX_IMAGE_PIXELS:
                        return {"state": "failed", "reason": "decoded_image_too_large"}
                    decoded.load()
            except Exception as exc:
                return {"state": "failed", "reason": "invalid_image_bytes", "error_type": type(exc).__name__}
            return {"state": "verified", "reason": "public_bytes_match", "bytes": size}
    except httpx.HTTPError as exc:
        # Exception text can contain a query or proxy credentials. Record type only.
        return {"state": "unverified", "reason": "request_failed", "error_type": type(exc).__name__}


def inspect_assets(
    value: Any,
    *,
    allow_probe: bool | None = None,
    client: httpx.Client | None = None,
) -> dict[str, Any]:
    """Inspect once per staging invocation; identical content hashes probe once.

    Failures remain report data so artifact generation always continues. A dry
    run never initiates network I/O by default. Tests may inject a mock client.
    """
    urls = image_urls(value)
    public_origin = configured_origin()
    valid_origin = valid_public_origin(public_origin)
    enabled = config.use_live_generation() if allow_probe is None else allow_probe
    report: dict[str, Any] = {
        "version": VERSION,
        "checked_at": datetime.now(timezone.utc).isoformat(),
        "configured_origin": public_origin,
        "inputs_sha256": _input_sha(urls),
        "source_url_count": len(urls),
        "assets": [],
        "state": "no_assets" if not urls else "unverified",
        "probe_limit": MAX_PROBES,
        "byte_limit": MAX_ASSET_BYTES,
        "time_budget_seconds": TOTAL_SECONDS,
    }
    own_client = None
    deadline = time.monotonic() + TOTAL_SECONDS
    by_identity: dict[str, dict[str, Any]] = {}
    attempts = 0
    try:
        for url in urls:
            try:
                parsed = urlsplit(url)
                route = _ROUTE.fullmatch(parsed.path)
                same_origin = valid_origin and _origin(url) == _origin(public_origin)
            except ValueError:
                report["assets"].append({
                    "source_urls": [url], "state": "unverified",
                    "reason": "malformed_image_url", "local_state": "not_our_asset_route",
                })
                continue
            trusted = bool(route and same_origin and not parsed.username and not parsed.password)
            identity = route[2] if trusted and route else url
            if identity in by_identity:
                by_identity[identity]["source_urls"].append(url)
                continue
            entry: dict[str, Any] = {"source_urls": [url], "state": "unverified"}
            by_identity[identity] = entry
            report["assets"].append(entry)
            if route:
                entry["sha256"] = route[2].removesuffix(".jpg")
                entry["local_state"] = _local_status(route[2])
            else:
                entry["local_state"] = "not_our_asset_route"
            if not valid_origin:
                entry["reason"] = "public_origin_unconfigured_or_invalid"
            elif not trusted:
                entry["reason"] = "outside_configured_asset_origin_or_route"
            elif not enabled:
                entry["reason"] = "public_probe_not_run"
            elif attempts >= MAX_PROBES or time.monotonic() >= deadline:
                entry["reason"] = "probe_budget_exhausted"
            else:
                if client is None:
                    own_client = own_client or httpx.Client(
                        timeout=httpx.Timeout(REQUEST_SECONDS),
                        follow_redirects=False,
                        trust_env=False,
                    )
                    client = own_client
                attempts += 1
                # The signature is advisory on our route; never forward a
                # query supplied by content, nor follow a redirect elsewhere.
                target = public_origin + parsed.path
                entry["probed_url"] = target
                entry.update(_probe(client, target, entry["sha256"], deadline))
    except Exception as exc:  # an incomplete audit is data, never a lost export
        report["audit_error"] = type(exc).__name__
    finally:
        if own_client is not None:
            own_client.close()
    report["probe_count"] = attempts
    report["verified_count"] = sum(row["state"] == "verified" for row in report["assets"])
    if urls:
        states = {row["state"] for row in report["assets"]}
        report["state"] = (
            "verified" if states == {"verified"} and not report.get("audit_error")
            else "failed" if "failed" in states else "unverified"
        )
    return report


def readiness_defects(payload: Mapping[str, Any]) -> list[str]:
    """Read the frozen audit mechanically; never re-probe while rendering."""
    report = payload.get(REPORT_FIELD)
    if not isinstance(report, Mapping):
        return []  # legacy staged releases predate this explicit audit
    urls = image_urls(release_asset_input(payload))
    if report.get("inputs_sha256") != _input_sha(urls):
        return ["source_asset_publication_stale: exported image references changed after the delivery check"]
    if not urls:
        return []
    return report_defects(report)


def report_defects(report: Mapping[str, Any]) -> list[str]:
    if report.get("state") not in {"verified", "no_assets"}:
        return [
            "source_asset_publication_unverified: exported images are not all verified at the configured public origin; "
            "local pinning alone is not public delivery (see source_asset_publication report)"
        ]
    return []


def inspect_workbooks(
    workbooks: Mapping[str, bytes],
    *,
    allow_probe: bool | None = None,
    client: httpx.Client | None = None,
) -> dict[str, Any]:
    """Probe the image references in the actual serialized workbook cells."""
    from openpyxl import load_workbook

    values: list[Any] = []
    hashes = {name: hashlib.sha256(data).hexdigest() for name, data in workbooks.items()}
    try:
        for data in workbooks.values():
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                for sheet in workbook:
                    headers: tuple[Any, ...] = ()
                    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        values.extend(cell for cell in row if isinstance(cell, str))
                        if row_number == 2:
                            headers = row
                        if row_number > 2 and headers:
                            record = dict(zip(headers, row))
                            for field, cell in record.items():
                                if not isinstance(field, str) or "answer_type_" not in field or str(cell).strip().lower() != "image":
                                    continue
                                for candidate in (
                                    field.replace("answer_type_", "answer_content_"),
                                    field.replace("answer_type_", "answer_"),
                                    field.replace("answer_type_", "keyword_"),
                                ):
                                    if candidate in record:
                                        values.append({"answer_type": "Image", "answer_content": record[candidate]})
            finally:
                workbook.close()
        report = inspect_assets(values, allow_probe=allow_probe, client=client)
    except Exception as exc:
        report = {"version": VERSION, "state": "unverified", "audit_error": type(exc).__name__, "assets": []}
    return {**report, "workbook_sha256s": hashes}


def release_findings(report: Mapping[str, Any]) -> list[dict[str, Any]]:
    if report.get("state") in {"verified", "no_assets"}:
        return []
    return [{
        "code": "source_asset_publication_unverified",
        "phase": "release_assets",
        "severity": "error",
        "message": (
            "Public image delivery is not fully verified. The artifacts remain available; "
            "inspect the source_asset_publication report before database publication. "
            "A locally pinned crop does not prove that its public URL serves those bytes."
        ),
    }]
