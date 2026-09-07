"""Mechanical contract rules the owner's corrected outputs demonstrate.

Evidence base: the corrected English and Mathematics workbooks supplied on
2026-09-06 (11 chapters, with their source chapters and per-chapter error
logs). These tests pin only the TECHNICAL layer those files fix — cell
types, marker values, display formats and referential integrity — never a
semantic decision, which stays the model's under CLAUDE.md Rule 1.

Where a corrected file and the contract disagree the contract governs
(§2 precedence, §45 "calibration is not authority"), so the delimiter and
rubric-quantum defects those hand-edited files carry are deliberately NOT
pinned here.
"""
from __future__ import annotations

import io

import openpyxl
import pytest

from app import models
from app.bulk_import import assessment_workbook as workbook
from app.bulk_import import layouts, writer


CONCEPT_LAYOUT = writer.CONCEPT_FILE_LAYOUT_ID


def _objective_row(data: bytes) -> tuple[dict, object]:
    """The first data row of a Concept File, by field name, plus its sheet."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheet = wb[writer.SHEET_BY_KIND["objective"]]
    header = [
        "" if cell.value is None else str(cell.value) for cell in sheet[2]
    ]
    values = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
    return {name: values[i] for i, name in enumerate(header) if name}, sheet


def _chapter_with_one_concept(db) -> models.Concept:
    chapter = models.Chapter(
        chapter_code="TECH", board="Maharashtra", grade="06",
        subject="English", unit="U", chapter_title="A Water Drop!",
        chapter_display_name="A Water Drop!",
        chapter_duration="200 minutes",
    )
    db.add(chapter)
    db.flush()
    topic = models.Topic(
        chapter_id=chapter.id, topic_title="Detailed Analysis",
        topic_display_name="Detailed Analysis", pre_post_learning="Post",
        source_order=1, topic_description="What the analysis covers.",
    )
    db.add(topic)
    db.flush()
    concept = models.Concept(
        topic_id=topic.id, concept_title="Language and Literary Devices",
        concept_display_name="Language and Literary Devices",
        concept_details="Description: d", source_order=1,
        sources="Balbharati",
    )
    db.add(concept)
    db.commit()
    return concept


# --------------------------------------------------------------------------- #
# §14.1 — the five update markers, on the Concept File too
# --------------------------------------------------------------------------- #

def test_concept_file_carries_all_five_update_markers(db):
    """Contract §14.1 / §42 gate 2, and the owner's corrected Concept File.

    [measured on "Corrected copy of 03_06_MSBSHSE_English_Radhas_Letter_to_
    Mowgli_Post_Concept.xlsm"] every one of the five reads ``No`` on a row
    whose Group and Question bands are empty. Before this contract slice
    Outputs 01/03 shipped ``is_update_group`` and ``is_update_question``
    blank, because the row builder padded everything past the Concept band.
    """
    concept = _chapter_with_one_concept(db)
    data = writer.write_concepts_workbook(
        db, [concept.id], layout_id=CONCEPT_LAYOUT, publication="Balbharati",
    )
    row, _sheet = _objective_row(data)
    assert [row[field] for field in workbook.UPDATE_FIELDS] == (
        [workbook.UPDATE_FIELD_VALUE] * 5
    )
    # The rest of the Group and Question bands stays empty: the markers are
    # the only cells that tail carries (§12 "hierarchy-only rows").
    tail = layouts.sheet(CONCEPT_LAYOUT, "objective").fields[
        len(layouts.sheet(CONCEPT_LAYOUT, "objective").block_fields("chapter"))
        + len(layouts.sheet(CONCEPT_LAYOUT, "objective").block_fields("topic"))
        + len(layouts.sheet(CONCEPT_LAYOUT, "objective").block_fields("concept")):
    ]
    for field in tail:
        if field in workbook.UPDATE_FIELDS:
            continue
        assert not str(row[field] or "").strip(), field


def test_concept_file_read_back_names_a_blank_update_marker(db):
    """The Concept File gets the Master's §42 gate-2 assertion as well."""
    concept = _chapter_with_one_concept(db)
    data = writer.write_concepts_workbook(
        db, [concept.id], layout_id=CONCEPT_LAYOUT, publication="Balbharati",
    )
    wb = openpyxl.load_workbook(io.BytesIO(data))
    sheet = wb[writer.SHEET_BY_KIND["objective"]]
    header = [
        "" if cell.value is None else str(cell.value) for cell in sheet[2]
    ]
    sheet.cell(row=3, column=header.index("is_update_group") + 1).value = None
    damaged = io.BytesIO()
    wb.save(damaged)

    decisions = writer._validate_concepts_workbook_bytes(
        damaged.getvalue(),
        [concept],
        writer.ConceptExportScope([concept]),
        exact_rows=True,
        sheet_layout=layouts.sheet(CONCEPT_LAYOUT, "objective"),
    )
    findings = [f for d in decisions for f in d["findings"]]
    assert any("is_update_group" in f for f in findings), findings


def test_concept_file_read_back_addresses_the_layout_it_wrote(db):
    """§14: validations locate fields by row-2 header name.

    The gate used to fall back to the 67-column reference layout while the
    bytes were written on the 72-column update-aware one, so it compared
    shifted columns and could neither confirm nor refute the topology.
    """
    concept = _chapter_with_one_concept(db)
    data = writer.write_concepts_workbook(
        db, [concept.id], layout_id=CONCEPT_LAYOUT, publication="Balbharati",
    )
    assert writer._validate_concepts_workbook_bytes(
        data,
        [concept],
        writer.ConceptExportScope([concept]),
        exact_rows=True,
        sheet_layout=layouts.sheet(CONCEPT_LAYOUT, "objective"),
    ) == []


# --------------------------------------------------------------------------- #
# §32 — real numeric cells with the 0.## display
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize(
    "text, expected",
    [("1", 1), ("1.0", 1.0), ("0", 0), ("0.5", 0.5), ("2.25", 2.25)],
)
def test_numeric_text_becomes_a_real_number(text, expected):
    value = workbook._numeric_cell(text)
    assert value == expected
    assert isinstance(value, (int, float))
    # The written text decides int or float, so nothing but the TYPE moves.
    assert isinstance(value, float) is isinstance(expected, float)


@pytest.mark.parametrize("value", ["", "   ", "two", "5 minutes", None])
def test_non_numeric_values_are_left_for_the_gates_to_name(value):
    """A value that is not exactly a number is never guessed at."""
    assert workbook._numeric_cell(value) is value


def test_concept_file_chapter_duration_is_numeric_with_the_contract_format(db):
    """§32: chapter_duration is a real numeric cell displayed as ``0.##``.

    The Master has carried the format since Q26; the Concept File set no
    number format at all, so the same value displayed differently in two
    of the four outputs.
    """
    concept = _chapter_with_one_concept(db)
    data = writer.write_concepts_workbook(
        db, [concept.id], layout_id=CONCEPT_LAYOUT, publication="Balbharati",
    )
    row, sheet = _objective_row(data)
    assert row["chapter_duration"] == 200
    header = [
        "" if cell.value is None else str(cell.value) for cell in sheet[2]
    ]
    cell = sheet.cell(row=3, column=header.index("chapter_duration") + 1)
    assert cell.number_format == workbook.NUMERIC_DISPLAY_FORMAT


def test_numeric_display_fields_cover_the_contract_list():
    """§32 names them: durations, marks and every weight."""
    for field in (
        "chapter_duration", "question_duration", "marks",
        "answer_weightage_1", "weightage_20",
        "sub_question_marks_15", "sq15_weightage_6",
    ):
        assert workbook.is_numeric_display_field(field), field
    for field in ("question", "concept_title", "group_name", "placeholder_1"):
        assert not workbook.is_numeric_display_field(field), field


# --------------------------------------------------------------------------- #
# §13 / §42 gate 4 — every parent list resolves inside the file
# --------------------------------------------------------------------------- #

def _parsed(rows_by_sheet: dict[str, list[dict]]) -> dict:
    return {
        "sheets": {
            name: {"rows": rows_by_sheet.get(name, []),
                   "row_numbers": list(
                       range(3, 3 + len(rows_by_sheet.get(name, []))))}
            for name in ("Objective", "Subjective", "Descriptive")
        }
    }


def test_a_label_naming_no_question_in_the_file_is_a_finding():
    """The owner's Love for One's Motherland log records this exact defect:
    a concept kept listing ``… Q01`` after that question was removed, so
    "the workbook now points to a question label that no longer exists
    anywhere in the file"."""
    parsed = _parsed({
        "Descriptive": [
            {
                "question_label": "06MSEN_X_PL_T01_C01 Q02",
                "concept_question_labels": (
                    "06MSEN_X_PL_T01_C01 Q01 | 06MSEN_X_PL_T01_C01 Q02"
                ),
                "group_question_labels": "06MSEN_X_PL_T01_C01 Q02",
                "group_name": "(06MSEN_X_PL_T01_C01) BG01",
                "basic_groups": "(06MSEN_X_PL_T01_C01) BG01",
            },
        ],
    })
    errors = workbook._dangling_reference_errors(parsed)
    assert len(errors) == 1
    assert "Q01" in errors[0]
    assert "concept_question_labels" in errors[0]


def test_a_group_rollup_naming_no_group_in_the_file_is_a_finding():
    parsed = _parsed({
        "Objective": [
            {
                "question_label": "06MSEN_X_PL_T01_C01 Q01",
                "group_name": "(06MSEN_X_PL_T01_C01) BG01",
                "concept_question_labels": "06MSEN_X_PL_T01_C01 Q01",
                "group_question_labels": "06MSEN_X_PL_T01_C01 Q01",
                "advanced_groups": "(06MSEN_X_PL_T01_C01) AG01",
            },
        ],
    })
    errors = workbook._dangling_reference_errors(parsed)
    assert len(errors) == 1
    assert "AG01" in errors[0]
    assert "advanced_groups" in errors[0]


def test_labels_resolving_across_sheets_are_not_dangling():
    """A concept's roster spans the whole file, not one sheet."""
    parsed = _parsed({
        "Objective": [
            {
                "question_label": "06MSEN_X_PL_T01_C01 Q01",
                "group_name": "(06MSEN_X_PL_T01_C01) BG01",
                "concept_question_labels": (
                    "06MSEN_X_PL_T01_C01 Q01 | 06MSEN_X_PL_T01_C01 Q02"
                ),
                "basic_groups": "(06MSEN_X_PL_T01_C01) BG01",
            },
        ],
        "Descriptive": [
            {
                "question_label": "06MSEN_X_PL_T01_C01 Q02",
                "group_name": "(06MSEN_X_PL_T01_C01) BG01",
            },
        ],
    })
    assert workbook._dangling_reference_errors(parsed) == []


def test_a_concept_only_tail_row_carries_no_reference():
    """The questionless tail row lists nothing, so it dangles nothing."""
    parsed = _parsed({
        "Objective": [
            {
                "question_label": "", "group_name": "",
                "concept_question_labels": "", "basic_groups": "",
            },
        ],
    })
    assert workbook._dangling_reference_errors(parsed) == []
