"""A new Post correction owns mini routing independently of its old source."""
from __future__ import annotations

import copy
import io

import pytest
from openpyxl import load_workbook

from app import models
from app.services import (
    assessment_release_run as master_run,
    assessment_release_snapshot as snapshot,
    build_concepts_release as release,
    build_concepts_release_files as release_files,
    concept_question_review as review,
    generation,
    generation_repair_policy as repair,
    model_provider,
    model_routing_run,
    release_workbook_edits as edits,
)
from tests import test_canonical_concept_review_handoff as canonical
from tests import test_pre_release_lane_wiring as pre_fixtures


def _corrected_workbook(db, job, path):
    workbook = load_workbook(io.BytesIO(
        release_files.build_release_bulk_import_workbook(db, job, lane="post")
    ))
    sheet, row, column = canonical._concept_details_cell(workbook)
    sheet.cell(row, column).value = (
        "Description: Supported core with one disputed clause. "
        "// Types: Type 01: Evidence-based explanation "
        "Case 01: Explain the foundational relationship. "
        "Case 02: Apply the reusable method later. "
        "Example 01: Explain the FIRST [review note omitted] method now. "
        "Example 02: Describe the manually added method. "
        "Manual context. Manual answer. Manual option A. Manual option B."
    )
    workbook.save(path)


def test_corrected_post_author_critic_and_master_use_mini_without_replacing_source_or_pre(db, monkeypatch, tmp_path):
    job = canonical._canonical_release(db)
    original_post = copy.deepcopy(release.release_payload(job, lane="post"))
    release.stage_pre_release(
        db, job, target_chapter_id=original_post["target_chapter_id"],
        pre_map=pre_fixtures._pre_map(), pre_questions=pre_fixtures._pre_questions(), inventory={},
    )
    original_pre = copy.deepcopy(release.release_payload(job, lane="pre"))
    model_routing_run.save_profile_for_job(job, model_provider.legacy_profile())
    routing_path = model_routing_run._record_path(job)
    original_routing = routing_path.read_bytes()
    path = tmp_path / "corrected-post.xlsx"
    _corrected_workbook(db, job, path)
    calls = []

    def provider(system, user, **kwargs):
        profile = model_provider.bound_profile()
        calls.append((kwargs["stage"], copy.deepcopy(profile)))
        assert profile == model_provider.new_profile()
        assert {route["model"] for route in profile["routes"].values()} == {"gpt-5.6-luna"}
        if kwargs["stage"] == "concept_review.critic":
            return {"verdict": "verified", "issues": []}
        return canonical._author_verdict()

    monkeypatch.setattr(generation, "_openai_json", provider)
    with model_routing_run.bind_job(job):
        assert model_provider.bound_profile() == model_provider.legacy_profile()
        result = edits.apply_workbook_for_review(
            db, job, lane="post", workbook_path=path, owner_sub=job.owner_sub,
        )
        assert model_provider.bound_profile() == model_provider.legacy_profile()
    assert result["round_recorded"] is True
    assert [stage for stage, _ in calls] == ["concept_review.author", "concept_review.critic"]
    accepted = release.release_payload(job, lane="post")
    assert accepted[model_provider.PROFILE_KEY] == model_provider.new_profile()
    assert repair.active(accepted)
    assert release.release_payload(job, lane="pre") == original_pre
    assert routing_path.read_bytes() == original_routing
    original_version = db.query(models.ConceptReleaseVersion).filter_by(
        job_id=job.id, lane="post", origin="staged",
    ).order_by(models.ConceptReleaseVersion.id).first()
    assert original_version.payload == original_post
    assert model_provider.PROFILE_KEY not in snapshot.build(db, job, original_pre)["metadata"]

    class MasterBoundaryReached(Exception):
        pass

    real_snapshot = snapshot.build
    master_calls = []

    def inspect_master(db, source_job, staged):
        bridge = real_snapshot(db, source_job, staged)
        profile = model_provider.bound_profile()
        master_calls.append((copy.deepcopy(profile), bridge["metadata"]))
        assert profile == model_provider.new_profile()
        assert bridge["metadata"][model_provider.PROFILE_KEY] == profile
        assert {route["provider"] for route in profile["routes"].values()} == {"openai"}
        assert {route["model"] for route in profile["routes"].values()} == {"gpt-5.6-luna"}
        raise MasterBoundaryReached

    monkeypatch.setattr(snapshot, "build", inspect_master)
    with model_routing_run.bind_job(job):
        with pytest.raises(MasterBoundaryReached):
            master_run.run_release_for_job(db, job.id, lane="post", owner_sub=job.owner_sub)
        assert model_provider.bound_profile() == model_provider.legacy_profile()
    assert len(master_calls) == 1
    assert routing_path.read_bytes() == original_routing
    assert release.release_payload(job, lane="pre") == original_pre


def test_unchanged_post_upload_does_not_create_revision_profile_or_call_provider(db, monkeypatch, tmp_path):
    job = canonical._canonical_release(db)
    before = copy.deepcopy(release.release_payload(job, lane="post"))
    model_routing_run.save_profile_for_job(job, model_provider.legacy_profile())
    path = tmp_path / "unchanged-post.xlsx"
    path.write_bytes(release_files.build_release_bulk_import_workbook(db, job, lane="post"))
    version_count = db.query(models.ConceptReleaseVersion).filter_by(job_id=job.id).count()
    monkeypatch.setattr(generation, "_openai_json", lambda *a, **k: pytest.fail("unchanged upload spent on review"))
    with model_routing_run.bind_job(job):
        result = edits.apply_workbook_for_review(
            db, job, lane="post", workbook_path=path, owner_sub=job.owner_sub,
        )
    assert result["round_recorded"] is False
    assert release.release_payload(job, lane="post") == before
    assert db.query(models.ConceptReleaseVersion).filter_by(job_id=job.id).count() == version_count
    assert model_routing_run.recorded_profile_for_job(job) == model_provider.legacy_profile()


def test_historical_v4_review_retains_its_inherited_profile(monkeypatch):
    profiles = []

    def provider(*args, **kwargs):
        profiles.append(model_provider.bound_profile())
        assert kwargs["response_schema"].name.endswith("_v4")
        return {}

    monkeypatch.setattr(generation, "_openai_json", provider)
    with model_provider.bind_profile(model_provider.legacy_profile()):
        review._call("review", {}, critic=False)
        review._call("review", {}, critic=True)
    assert profiles == [model_provider.legacy_profile(), model_provider.legacy_profile()]
