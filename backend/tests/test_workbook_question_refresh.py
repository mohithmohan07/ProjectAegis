"""Q51 D14 — the opt-in Question-band refresh of the shared CMS workbook.

``append_questions`` is append-only, and stays append-only for every caller
that does not NAME the labels it may rewrite. The owner approved one explicit
exception so a question re-published with edited content stops carrying its
superseded wording in the export until somebody re-exports by hand
(``docs/three-step-workflow-review-2026-09-11.md`` §7 D14).

These pin the shape of that exception:

* a caller that names nothing cannot overwrite a single cell (the default);
* a named label's row is re-projected IN PLACE — same row, same position,
  Chapter/Topic/Concept and Group bands untouched, no row added or removed,
  and no other label's row touched;
* repeating the write converges instead of double-writing.
"""

from pathlib import Path

import openpyxl

from app import bulk_import as bi
from app import models
from app.bulk_import import writer


def _graph(db, suffix: str):
    """One chapter with two Objective questions under one concept."""
    chapter = models.Chapter(
        chapter_code=f"10CBMA_QRefresh_{suffix}",
        board="CBSE",
        grade="10",
        subject="Mathematics",
        unit="Question Refresh",
        chapter_title=f"Question Refresh {suffix}",
        chapter_display_name=f"Question Refresh {suffix}",
    )
    topic = models.Topic(
        chapter=chapter,
        topic_title="Home Topic",
        topic_display_name="Home Topic",
        pre_post_learning="Post",
    )
    concept = models.Concept(
        topic=topic,
        concept_title="Ohm's Law",
        concept_display_name="Ohm's Law",
        concept_details="Original concept details",
    )
    group = models.Group(
        concept=concept,
        group_type="Basic",
        group_name=f"({suffix}) BG01",
        group_display_name=f"({suffix}) BG01",
        group_status="Active",
    )
    first = models.Question(
        group=group,
        sheet_kind="objective",
        origin="concept_mapping",
        question_label=f"QREFRESH-{suffix} Q01",
        question_category="Multiple Choice Question",
        cognitive_skills="Remember",
        question_source="NCERT",
        question="Which statement describes Ohm's law?",
        question_text="Which statement describes Ohm's law?",
        level_of_difficulty="Less",
        marks=1,
        answers=[],
    )
    second = models.Question(
        group=group,
        sheet_kind="objective",
        origin="concept_mapping",
        question_label=f"QREFRESH-{suffix} Q02",
        question_category="Multiple Choice Question",
        cognitive_skills="Remember",
        question_source="NCERT",
        question="Which unit measures resistance?",
        question_text="Which unit measures resistance?",
        level_of_difficulty="Less",
        marks=1,
        answers=[],
    )
    db.add(chapter)
    db.commit()
    return concept, group, first, second


def _rows(path: Path) -> dict[int, tuple]:
    worksheet = openpyxl.load_workbook(path)[bi.SHEET_OBJECTIVE]
    return {
        row_i: tuple(cell.value for cell in worksheet[row_i])
        for row_i in range(3, worksheet.max_row + 1)
        if any(cell.value is not None for cell in worksheet[row_i])
    }


def test_a_caller_that_names_no_labels_cannot_overwrite_an_existing_row(
    db, tmp_path,
):
    """The default is still strictly append-only (post_generation's contract)."""
    _concept, _group, first, second = _graph(db, "Default")
    path = tmp_path / "output.xlsx"
    assert writer.append_questions(db, path, [first.id, second.id])["objective"] == 2
    before = _rows(path)

    first.question = "Rewritten in the database only."
    first.question_text = "Rewritten in the database only."
    first.level_of_difficulty = "Moderate"
    db.commit()

    receipt = writer.append_questions(db, path, [first.id, second.id])

    assert receipt["objective"] == 0
    assert receipt["skipped"] == 2
    # Nothing about a refresh is even reported to a caller that asked for none.
    assert "refreshed" not in receipt
    assert "refreshed_labels" not in receipt
    # …and not one cell moved.
    assert _rows(path) == before


def test_a_named_label_is_refreshed_in_place_and_no_other_row_is_touched(
    db, tmp_path,
):
    _concept, _group, first, second = _graph(db, "Named")
    path = tmp_path / "output.xlsx"
    assert writer.append_questions(db, path, [first.id, second.id])["objective"] == 2
    before = _rows(path)
    columns = {
        str(cell.value): cell.column - 1
        for cell in openpyxl.load_workbook(path)[bi.SHEET_OBJECTIVE][2]
        if cell.value is not None
    }
    edited_row = next(
        row_i for row_i, values in before.items()
        if values[columns["question_label"]] == first.question_label
    )
    untouched_row = next(
        row_i for row_i, values in before.items()
        if values[columns["question_label"]] == second.question_label
    )

    first.question = "Which statement describes Ohm's law exactly?"
    first.question_text = "Which statement describes Ohm's law exactly?"
    first.level_of_difficulty = "Moderate"
    db.commit()

    receipt = writer.append_questions(
        db, path, [first.id, second.id],
        refresh_labels=[first.question_label],
    )

    assert receipt["refreshed"] == 1
    assert receipt["refreshed_labels"] == [first.question_label]
    assert receipt["refreshed_unchanged"] == 0
    # The other placement is still an ordinary append-only skip.
    assert receipt["objective"] == 0
    assert receipt["skipped"] == 1
    assert receipt["skipped_reasons"] == {"already_present": 1}

    after = _rows(path)
    # Same rows, same positions, nothing added or removed.
    assert set(after) == set(before)
    assert after[untouched_row] == before[untouched_row]
    assert after[edited_row][columns["question"]] == (
        "Which statement describes Ohm's law exactly?")
    assert after[edited_row][columns["level_of_difficulty"]] == "Moderate"
    # Row identity and every band outside the Question band are untouched.
    assert after[edited_row][columns["question_label"]] == first.question_label
    for field in (
        "chapter_code", "topic_title", "concept_title", "group_name",
    ):
        if field in columns:
            assert after[edited_row][columns[field]] == (
                before[edited_row][columns[field]])


def test_repeating_the_refresh_converges_instead_of_writing_twice(db, tmp_path):
    _concept, _group, first, second = _graph(db, "Converge")
    path = tmp_path / "output.xlsx"
    writer.append_questions(db, path, [first.id, second.id])

    first.question = "Converged wording."
    first.question_text = "Converged wording."
    db.commit()

    once = writer.append_questions(
        db, path, [first.id, second.id],
        refresh_labels=[first.question_label],
    )
    settled = _rows(path)
    assert once["refreshed"] == 1

    twice = writer.append_questions(
        db, path, [first.id, second.id],
        refresh_labels=[first.question_label],
    )

    assert twice["refreshed"] == 0
    assert twice["refreshed_unchanged"] == 1
    assert twice["refreshed_unchanged_labels"] == [first.question_label]
    assert twice["objective"] == 0
    assert _rows(path) == settled


def test_an_unnamed_label_is_never_refreshed_even_beside_a_named_one(
    db, tmp_path,
):
    """Naming one label is not naming them all — no "refresh everything"."""
    _concept, _group, first, second = _graph(db, "OnlyNamed")
    path = tmp_path / "output.xlsx"
    writer.append_questions(db, path, [first.id, second.id])
    before = _rows(path)

    first.question = "Named and rewritten."
    first.question_text = "Named and rewritten."
    second.question = "Not named, so not rewritten."
    second.question_text = "Not named, so not rewritten."
    db.commit()

    receipt = writer.append_questions(
        db, path, [first.id, second.id],
        refresh_labels=[first.question_label],
    )

    assert receipt["refreshed"] == 1
    assert receipt["refresh_labels_requested"] == [first.question_label]
    after = _rows(path)
    columns = {
        str(cell.value): cell.column - 1
        for cell in openpyxl.load_workbook(path)[bi.SHEET_OBJECTIVE][2]
        if cell.value is not None
    }
    stale = next(
        row_i for row_i, values in before.items()
        if values[columns["question_label"]] == second.question_label
    )
    assert after[stale] == before[stale]
