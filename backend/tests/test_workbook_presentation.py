"""Real XLSX display/read-back regressions for readable, lossless rich cells."""
from __future__ import annotations

import copy
import io

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app import bulk_import as bi, models
from app.bulk_import import assessment_workbook as assessment
from app.bulk_import import layouts, reader, writer
from app.bulk_import.presentation import to_display_rich_text
from tests.test_mes_dual_output import _snapshot
from tests.test_katex_render_validation import real_engine


_IMAGE = '[img src="https://projectaegis.fly.dev/assets/chapter/figure-01.png" alt="Source figure"]'
_RICH_TEXT = (
    "Description: Read the evidence.\n\n"
    "| Object | Length |\n| --- | --- |\n| Rod | [Katex]2x[/Katex] |\n"
    + _IMAGE + "\nUse the figure and both columns."
)
_MULTILINE_TEX = (
    r"\begin{array}{|l|r|}\hline" "\n"
    "% The following data row must remain outside this comment.\n"
    r"\text{Rod} & 2x \\ \hline" "\n"
    r"\text{Wire} & x+1 \\ \hline" "\n"
    r"\end{array}"
)
_WRAPPED_TEX = "[Katex]" + _MULTILINE_TEX + "[/Katex]"


def _equation_snapshot(kind: str, tex: str = _MULTILINE_TEX):
    """Explicitly declared Equation cells in every supported answer slot."""
    snapshot = _snapshot()
    candidate = copy.deepcopy(snapshot["candidates"][0 if kind in {"objective", "subjective"} else 1])
    candidate["sheet_kind"] = "descriptive" if kind == "child" else kind
    candidate["sub_questions"] = []
    candidate["answers"] = [{
        "answer_type": "Equation", "answer_content": tex,
        "correct_answer": "1", "answer_weightage": "1", "placeholder": "a",
    }]
    if kind == "subjective":
        candidate["question"] = candidate["question_text"] = "Complete $$a$$."
    elif kind == "child":
        candidate["answers"] = []
        candidate["sub_questions"] = [{
            "text": "Write the table.", "marks": "1",
            "keywords": [{"answer_type": "Equation", "keyword": tex, "weightage": "1"}],
        }]
    snapshot["candidates"] = [candidate]
    field = "sq1_keyword_1" if kind == "child" else "answer_1" if kind == "subjective" else "answer_content_1"
    return snapshot, candidate["sheet_kind"].capitalize(), field


def _open(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


def _field_cell(ws, field: str, row: int = 3):
    fields = tuple(cell.value for cell in ws[2])
    return ws.cell(row=row, column=fields.index(field) + 1)


@pytest.mark.parametrize(
    ("source", "display", "model"),
    [
        ("first\nsecond", "first<br>\nsecond", "first\nsecond"),
        ("first\n\nsecond", "first<br>\n<br>\nsecond", "first\n\nsecond"),
        ("first<br>second", "first<br>\nsecond", "first\nsecond"),
        ("first<br><br>second", "first<br>\n<br>\nsecond", "first\n\nsecond"),
        ("first<BR />\r\nsecond", "first<br>\nsecond", "first\nsecond"),
        ("first\r\n\r\nsecond", "first<br>\n<br>\nsecond", "first\n\nsecond"),
        ("first<br>\n<br>\nsecond", "first<br>\n<br>\nsecond", "first\n\nsecond"),
        ("first<br>\n\nsecond", "first<br>\n<br>\nsecond", "first\n\nsecond"),
    ],
)
def test_display_breaks_do_not_multiply_on_import_or_reexport(source, display, model):
    assert to_display_rich_text(source) == display
    assert bi.from_workbook_rich_text(display) == model
    assert to_display_rich_text(display) == display
    assert to_display_rich_text(bi.from_workbook_rich_text(display)) == display


def test_canonical_projection_remains_distinct_from_excel_presentation():
    assert bi.to_workbook_rich_text("first\n\nsecond") == "first<br><br>second"
    assert bi.from_workbook_rich_text("first<br><br>second") == "first\n\nsecond"


def test_assessment_readback_keeps_canonical_wire_records():
    snapshot = _snapshot()
    snapshot["topics"][0]["concepts"][0]["concept_details"] = "First\n\nSecond"
    data = assessment.render_concept_file(snapshot)
    row = assessment.parse_workbook(data)["sheets"]["Objective"]["rows"][0]
    assert row["concept_details"] == "First<br><br>Second"
    assert bi.from_workbook_rich_text(row["concept_details"]) == "First\n\nSecond"
    assert row["chapter_duration"] == 40


def test_public_writer_keeps_tables_figures_and_paragraphs_readable_on_readback(db):
    question = db.query(models.Question).order_by(models.Question.id).first()
    assert question is not None
    original = question.question
    try:
        question.question = _RICH_TEXT
        db.flush()
        output = writer.write_workbook(db, question_ids=[question.id])
        wb = _open(output)
        ws = wb[bi.SHEET_BY_KIND[question.sheet_kind]]
        cell = _field_cell(ws, "question")
        # This legacy database exporter already projects Markdown tables to
        # supported KaTeX arrays. Presentation must preserve its whole result.
        expected_model = (
            "Description: Read the evidence.\n\n"
            r"[Katex] \begin{array}{|l|l|} \hline \text{Object} & \text{Length} "
            r"\\ \hline \text{Rod} & 2x \\ \hline \end{array} [/Katex]"
            "\n" + _IMAGE + "\nUse the figure and both columns."
        )
        expected = expected_model.replace("\n", "<br>\n")
        assert cell.value == expected
        assert _IMAGE in cell.value
        assert r"\text{Rod} & 2x" in cell.value
        assert cell.alignment.wrap_text is True
        assert cell.alignment.vertical == "top"
        found = layouts.identify_sheet(ws.title, tuple(c.value for c in ws[2]))
        assert found is not None
        readback = reader._block(found.sheet, tuple(c.value for c in ws[3]), "question")
        assert readback["question"] == expected_model
        question.question = readback["question"]
        second = _open(writer.write_workbook(db, question_ids=[question.id]))
        assert _field_cell(second[ws.title], "question").value == expected
        second.close()
        wb.close()
    finally:
        question.question = original
        db.rollback()


@pytest.mark.parametrize("role", ["concept", "master"])
def test_release_workbook_preserves_header_geometry_and_styles_rich_fields(role):
    snapshot = _snapshot()
    snapshot["topics"][0]["concepts"][0]["concept_details"] = _RICH_TEXT
    data = (
        assessment.render_concept_file(snapshot)
        if role == "concept"
        else assessment.render_master_file(snapshot)[0]
    )
    wb = _open(data)
    schema = assessment.output_schema(role, snapshot=snapshot)
    assert wb.sheetnames == assessment.SHEET_ORDER
    for ws in wb.worksheets:
        assert [c.value for c in ws[2]] == schema["fields"][ws.title]
        assert ws.max_column == len(schema["fields"][ws.title])
        assert ws.freeze_panes == "A3"
        expected_bands = {
            f"{get_column_letter(band['start'])}1:{get_column_letter(band['end'])}1"
            for band in schema["bands"][ws.title]
            if band["end"] > band["start"]
        }
        assert {str(region) for region in ws.merged_cells.ranges} == expected_bands
        fields = schema["fields"][ws.title]
        rich_col = get_column_letter(fields.index("concept_details") + 1)
        numeric_col = get_column_letter(fields.index("chapter_duration") + 1)
        assert ws.column_dimensions[rich_col].width > ws.column_dimensions[numeric_col].width
    cell = _field_cell(wb["Objective"], "concept_details")
    assert cell.value == _RICH_TEXT.replace("\n", "<br>\n")
    assert bi.from_workbook_rich_text(cell.value) == _RICH_TEXT
    assert cell.alignment.wrap_text is True
    assert cell.alignment.vertical == "top"
    # Excel is free to autosize the wrapped row; a fixed short row clips it.
    assert wb["Objective"].row_dimensions[3].height is None
    wb.close()


def test_master_options_and_multipart_children_have_real_line_breaks():
    snapshot = _snapshot()
    snapshot["candidates"][0]["question"] = "Look at the figure.\n\nWhich shape is flat?"
    snapshot["candidates"][0]["question_text"] = snapshot["candidates"][0]["question"]
    wb = _open(assessment.render_master_file(snapshot)[0])
    objective = _field_cell(wb["Objective"], "question_text").value
    assert objective == (
        "Look at the figure.<br>\n<br>\nWhich shape is flat?<br>\n"
        "a) Circle<br>\nb) Sphere"
    )
    descriptive = _field_cell(wb["Descriptive"], "question_text").value
    assert "<br>\n" in descriptive
    logical = bi.from_workbook_rich_text(descriptive)
    for part in snapshot["candidates"][1]["sub_questions"]:
        assert part["text"] in logical
    assert len(logical.splitlines()) >= 3
    wb.close()


@pytest.mark.parametrize("formula", ["=SUM(A1:A2)", "+2", "-3", "@A1"])
def test_display_preserves_literal_formula_text_and_real_numeric_cells(formula):
    snapshot = _snapshot()
    snapshot["topics"][0]["concepts"][0]["concept_details"] = formula + "\nExplanation"
    snapshot["candidates"][0]["marks"] = 1.5
    wb = _open(assessment.render_master_file(snapshot)[0])
    ws = wb["Objective"]
    formula_cell = _field_cell(ws, "concept_details")
    assert formula_cell.value == formula + "<br>\nExplanation"
    assert formula_cell.data_type == "s"
    number = _field_cell(ws, "marks")
    assert number.value == 1.5
    assert number.data_type == "n"
    assert number.number_format == "0.##"
    wb.close()


def test_ordinary_export_refuses_display_expansion_before_excel_can_truncate(db):
    question = db.query(models.Question).order_by(models.Question.id).first()
    assert question is not None
    original = question.question
    # Canonical br-only text fits, but the LF needed by Excel exceeds its cap.
    source = "x" * 32_762 + "\ny"
    assert len(bi.to_workbook_rich_text(source)) == writer.EXCEL_CELL_CHARACTER_LIMIT
    try:
        question.question = source
        db.flush()
        with pytest.raises(writer.ExcelCellLimitError, match="32,768 characters"):
            writer.write_workbook(db, question_ids=[question.id])
    finally:
        question.question = original
        db.rollback()


def test_staging_detects_the_same_display_expansion_as_the_cell_writer():
    from app.services import assessment_release

    source = "x" * (assessment.CELL_LIMIT - 4) + "\n"
    assert len(bi.to_workbook_rich_text(source)) == assessment.CELL_LIMIT
    findings = assessment_release._cell_shape_findings(
        {"concept_details": source}, "presentation-test",
    )
    assert len(findings) == 1
    assert findings[0]["reason"] == assessment.CELL_TEXT_TOO_LONG
    assert findings[0]["actual"] == assessment.CELL_LIMIT + 1


@pytest.mark.parametrize("role", ["concept", "master"])
def test_release_retains_complete_original_when_display_expansion_exceeds_cap(role):
    snapshot = _snapshot()
    source = "x" * 32_762 + "\ny"
    assert len(bi.to_workbook_rich_text(source)) == assessment.CELL_LIMIT
    snapshot["topics"][0]["concepts"][0]["concept_details"] = source
    if role == "concept":
        oversized = []
        data = assessment.render_concept_file(snapshot, oversized=oversized)
    else:
        data, issues = assessment.render_master_file(snapshot)
        oversized = issues["oversized_cells"]
    relevant = [item for item in oversized if item["context"].endswith(":concept_details")]
    assert relevant
    assert all(item["full_value"] == source for item in relevant)
    assert all(item["actual"] == assessment.CELL_LIMIT + 1 for item in relevant)
    wb = _open(data)
    value = _field_cell(wb["Objective"], "concept_details").value
    assert len(value) <= assessment.CELL_LIMIT
    assert "complete value is recorded" in value
    assert "oversized_cells" in value
    assert "<br>\n[Aegis:" in value
    wb.close()


def test_complete_katex_spans_preserve_native_linefeeds_and_tex_comments():
    source = "Read the table.\n\n" + _WRAPPED_TEX + "\nExplain both rows."
    canonical = "Read the table.<br><br>" + _WRAPPED_TEX + "<br>Explain both rows."
    displayed = "Read the table.<br>\n<br>\n" + _WRAPPED_TEX + "<br>\nExplain both rows."
    assert bi.to_workbook_rich_text(source) == canonical
    assert to_display_rich_text(source) == displayed
    assert bi.from_workbook_rich_text(canonical) == source
    assert bi.from_workbook_rich_text(displayed) == source
    assert to_display_rich_text(displayed) == displayed


def test_public_writer_and_reader_preserve_multiline_katex_without_html_inside(db):
    question = db.query(models.Question).order_by(models.Question.id).first()
    assert question is not None
    original = question.question
    source = "Read the table.\n" + _WRAPPED_TEX + "\n\nExplain both rows."
    try:
        question.question = source
        db.flush()
        wb = _open(writer.write_workbook(db, question_ids=[question.id]))
        ws = wb[bi.SHEET_BY_KIND[question.sheet_kind]]
        assert _field_cell(ws, "question").value == (
            "Read the table.<br>\n" + _WRAPPED_TEX + "<br>\n<br>\nExplain both rows."
        )
        found = layouts.identify_sheet(ws.title, tuple(c.value for c in ws[2]))
        assert found is not None
        readback = reader._block(found.sheet, tuple(c.value for c in ws[3]), "question")
        assert readback["question"] == source
        wb.close()
    finally:
        question.question = original
        db.rollback()


@pytest.mark.parametrize("role", ["concept", "master"])
def test_assessment_export_and_parser_preserve_multiline_katex_spans(role):
    snapshot = _snapshot()
    source = "Read the table.\n\n" + _WRAPPED_TEX + "\nExplain both rows."
    snapshot["topics"][0]["concepts"][0]["concept_details"] = source
    data = assessment.render_concept_file(snapshot) if role == "concept" else assessment.render_master_file(snapshot)[0]
    wb = _open(data)
    assert _field_cell(wb["Objective"], "concept_details").value == (
        "Read the table.<br>\n<br>\n" + _WRAPPED_TEX + "<br>\nExplain both rows."
    )
    canonical = assessment.parse_workbook(data)["sheets"]["Objective"]["rows"][0]["concept_details"]
    assert canonical == "Read the table.<br><br>" + _WRAPPED_TEX + "<br>Explain both rows."
    assert bi.from_workbook_rich_text(canonical) == source
    wb.close()


@pytest.mark.parametrize("kind", ["objective", "descriptive", "subjective", "child"])
def test_typed_equation_cells_keep_native_tex_in_every_assessment_slot(kind):
    snapshot, sheet, field = _equation_snapshot(kind)
    data, _ = assessment.render_master_file(snapshot)
    wb = _open(data)
    assert _field_cell(wb[sheet], field).value == _MULTILINE_TEX
    parsed = assessment.parse_workbook(data)["sheets"][sheet]["rows"][0]
    assert parsed[field] == _MULTILINE_TEX
    found = layouts.identify_sheet(sheet, tuple(c.value for c in wb[sheet][2]))
    assert found is not None
    imported = reader._block(found.sheet, tuple(c.value for c in wb[sheet][3]), "question")
    assert imported[field] == _MULTILINE_TEX
    wb.close()


@pytest.mark.parametrize("kind", ["objective", "descriptive", "subjective", "child"])
def test_public_database_export_preserves_all_typed_equation_slots(db, kind):
    question = db.query(models.Question).order_by(models.Question.id).first()
    assert question is not None
    snapshot, sheet, field = _equation_snapshot(kind)
    candidate = snapshot["candidates"][0]
    names = ("sheet_kind", "answers", "sub_questions", "question", "question_text", "marks")
    original = {name: copy.deepcopy(getattr(question, name)) for name in names}
    try:
        for name in names:
            setattr(question, name, copy.deepcopy(candidate[name]))
        db.flush()
        wb = _open(writer.write_workbook(db, question_ids=[question.id]))
        ws = wb[bi.SHEET_BY_KIND[candidate["sheet_kind"]]]
        assert _field_cell(ws, field).value == _MULTILINE_TEX
        found = layouts.identify_sheet(ws.title, tuple(c.value for c in ws[2]))
        assert found is not None
        imported = reader._block(found.sheet, tuple(c.value for c in ws[3]), "question")
        assert imported[field] == _MULTILINE_TEX
        wb.close()
    finally:
        for name, value in original.items():
            setattr(question, name, value)
        db.rollback()


def test_raw_equation_at_excel_limit_has_no_artificial_line_break_expansion(db):
    from app.services import assessment_release

    # Padding is a valid TeX comment; preserving LF keeps the equation alive.
    prefix, suffix = "x + %", "\n1"
    tex = prefix + "p" * (assessment.CELL_LIMIT - len(prefix) - len(suffix)) + suffix
    assert assessment_release._cell_shape_findings(
        {"answer_type_1": "Equation", "answer_content_1": tex}, "raw-equation-test",
    ) == []
    snapshot, sheet, field = _equation_snapshot("objective", tex)
    data, issues = assessment.render_master_file(snapshot)
    wb = _open(data)
    assert _field_cell(wb[sheet], field).value == tex
    assert not any(item["context"].endswith(":" + field) for item in issues["oversized_cells"])
    wb.close()
    question = db.query(models.Question).order_by(models.Question.id).first()
    assert question is not None
    original_kind, original_answers = question.sheet_kind, copy.deepcopy(question.answers)
    try:
        question.sheet_kind = "objective"
        question.answers = snapshot["candidates"][0]["answers"]
        db.flush()
        wb = _open(writer.write_workbook(db, question_ids=[question.id]))
        assert _field_cell(wb[bi.SHEET_BY_KIND["objective"]], field).value == tex
        wb.close()
    finally:
        question.sheet_kind, question.answers = original_kind, original_answers
        db.rollback()


@real_engine
def test_exported_multiline_math_and_comments_render_in_actual_katex_engine():
    from app.services import katex_render_validation as validation

    snapshot, sheet, field = _equation_snapshot("objective")
    snapshot["candidates"][0]["question"] = "Use this table.\n" + _WRAPPED_TEX
    data, _ = assessment.render_master_file(snapshot)
    parsed = assessment.parse_workbook(data)["sheets"][sheet]["rows"][0]
    report = validation.inspect_render({
        "question": parsed["question"],
        "answer_type_1": "Equation", "answer_content_1": parsed[field],
    })
    assert validation.report_defects(report) == []
    assert report["expression_count"] == 1
    assert report["occurrence_count"] == 2
    assert len(report["expressions"][0]["rendered_sha256"]) == 64
