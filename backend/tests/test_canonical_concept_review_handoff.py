"""Canonical three sheet Concept workbook -> reviewed Post bank handoff."""
from __future__ import annotations

import copy
import io
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

from app.services import assessment_source_inventory as source_inventory
from app.services import build_concepts_release as release
from app.services import build_concepts_release_files as release_files
from app.services import concept_question_review
from app.services import release_workbook_edits

sys.path.insert(0, str(Path(__file__).parent))
from test_build_concepts_release import (
    _inventory,
    _job,
    _mined_types,
    _rendered_records,
)


def _concept_details_cell(workbook):
    """Return the only canonical Concept Details cell and its sheet/column."""

    for sheet in workbook.worksheets:
        headers = {
            str(cell.value or "").strip().casefold(): cell.column
            for cell in sheet[2]
        }
        details_column = headers.get("concept_details")
        title_column = headers.get("concept_title")
        if not details_column or not title_column:
            continue
        for row in range(3, sheet.max_row + 1):
            if sheet.cell(row, title_column).value not in (None, ""):
                return sheet, row, details_column
    raise AssertionError("canonical workbook has no Concept Details row")


def _author_verdict(*, resolved=False):
    verdict = {
        "questions": [
            {
                "source_qid": "QINV-0001",
                "concept_row": 0,
                "placement_section": "types",
                "question_text": "Explain the FIRST method now.",
                "question_text_spans": ["Explain the FIRST", " method now."],
                "shared_context": "",
                "context_review": {
                    "action": "inherit",
                    "sources": [],
                    "rationale": "The original supporting context remains applicable.",
                },
                "source_answer": "",
                "options": [],
                "removed_dependencies": [],
                "type_id": "TYPE-0001",
                "type_title": "",
                "type_definition": "",
                "case_id": "CASE-0002",
                "case_definition": "",
                "rationale": "The retained question was moved to Case 02.",
            },
            {
                "source_qid": "",
                "concept_row": 0,
                "placement_section": "types",
                "question_text": "Describe the manually added method.",
                "question_text_spans": [],
                "shared_context": "Manual context.",
                "context_review": {
                    "action": "replace",
                    "sources": [{
                        "kind": "edited_concept",
                        "concept_row": 0,
                        "source_qid": "",
                        "reason": "The edited row supplies the added context.",
                    }],
                    "rationale": "The edited row supplies the context for the added task.",
                },
                "source_answer": "Manual answer.",
                "options": ["Manual option A", "Manual option B"],
                "removed_dependencies": [],
                "type_id": "TYPE-0001",
                "type_title": "",
                "type_definition": "",
                "case_id": "CASE-0002",
                "case_definition": "",
                "rationale": "The reviewer added this Example in Concept Details.",
            },
        ],
        "original_dispositions": [
            {
                "source_qid": "QINV-0001",
                "disposition": "moved",
                "rationale": "Retained with a reviewed Case target.",
            },
            {
                "source_qid": "QINV-0002",
                "disposition": "omitted",
                "rationale": "The reviewer removed Example 02.",
            },
        ],
        "row_dispositions": [{
            "concept_row": 0,
            "rationale": "The edited Types/Cases surface is authoritative.",
        }],
    }
    # A newly submitted Post correction uses the grouped v5 wire contract,
    # including explicit singleton/added membership and dependency choices.
    for question in verdict["questions"]:
        qid = question["source_qid"]
        question["source_qids"] = [qid] if qid else []
        question["source_dependency_reviews"] = [{
            "source_qid": qid, "removed_dependencies": [],
            "rationale": "The source dependencies remain applicable.",
        }] if qid else []
    if resolved:
        # The server resolves an explicit inherit decision before the critic
        # and before persisting the accepted author receipt.
        verdict["questions"][0]["shared_context"] = "Shared original context."
    return verdict


def _canonical_release(db):
    job, chapter = _job(db)
    records = _rendered_records()
    # The canonical writer is the source of the reviewer file's projected
    # parent value; this fixture keeps that value aligned for a true no-op
    # roundtrip while retaining the Type/Case release routes.
    records[0]["parent_concept"] = ""
    inventory = copy.deepcopy(_inventory())
    inventory["items"][0].update({
        "shared_context": "Shared original context.",
        "raw_solution_or_answer": "Original answer.",
        "options": ["Original option A", "Original option B"],
        "frozen_task_text": "Explain the first method.",
    })
    release.stage_release(
        db,
        job,
        target_chapter_id=chapter.id,
        records=records,
        inventory=inventory,
        mined_types=_mined_types(),
    )
    return job


def test_canonical_three_sheet_upload_reconciles_exact_reviewed_bank(
    db, tmp_path: Path, monkeypatch
):
    """Types/Cases edits in the downloaded canonical file reach the bank."""

    job = _canonical_release(db)
    workbook = load_workbook(io.BytesIO(
        release_files.build_release_bulk_import_workbook(db, job, lane="post")
    ))
    assert set(workbook.sheetnames) == {"Objective", "Descriptive", "Subjective"}
    sheet, row, details_column = _concept_details_cell(workbook)
    sheet.cell(row, details_column).value = (
        "Description: Supported core with one disputed clause. "
        "// Types: Type 01: Evidence-based explanation "
        "Case 01: Explain the foundational relationship. "
        "Case 02: Apply the reusable method later. "
        "Example 01: Explain the FIRST [review note omitted] method now. "
        "Example 02: Describe the manually added method. "
        "Manual context. Manual answer. Manual option A. Manual option B."
    )
    workbook_path = tmp_path / "canonical-corrected.xlsx"
    workbook.save(workbook_path)

    calls = []

    def fake_call(system, payload, *, critic=False):
        calls.append({"critic": critic, "payload": payload})
        if critic:
            return {"verdict": "verified", "issues": []}
        return _author_verdict()

    monkeypatch.setattr(concept_question_review, "_call", fake_call)
    result = release_workbook_edits.apply_workbook_for_review(
        db,
        job,
        lane="post",
        workbook_path=workbook_path,
        owner_sub="canonical-reviewer",
    )
    db.refresh(job)
    payload = release.release_payload(job, lane="post")
    assert payload is not None

    items = payload["question_task_inventory"]["items"]
    assert len(items) == 2
    assert items[0]["qid"] == "QINV-0001"
    assert items[0]["raw_task"] == "Explain the FIRST method now."
    assert items[0]["normalized_task"] == "Explain the FIRST method now."
    assert items[0]["polished_task"] == "Explain the FIRST method now."
    assert items[0]["frozen_task_text"] == "Explain the FIRST method now."
    assert items[0]["shared_context"] == "Shared original context."
    assert items[0]["raw_solution_or_answer"] == "Original answer."
    assert items[0]["options"] == ["Original option A", "Original option B"]
    assert items[0]["placement_section"] == "types"
    assert items[0]["_aegis_reviewed_target"]["case_id"] == "CASE-0002"

    added = items[1]
    assert added["qid"].startswith("QREV-")
    assert added["provenance"] == "reviewer_added"
    assert added["raw_task"] == "Describe the manually added method."
    assert added["shared_context"] == "Manual context."
    assert added["raw_solution_or_answer"] == "Manual answer."
    assert added["options"] == ["Manual option A", "Manual option B"]
    assert added["placement_section"] == "types"
    assert added["_aegis_reviewed_target"]["case_id"] == "CASE-0002"

    audit = payload["review_question_audit"]
    from app.services import generation_repair_policy as repair
    assert payload[repair.KEY] == repair.VERSION
    assert audit["model_review_receipt"]["policy"].startswith(concept_question_review.GROUP_POLICY)
    assert audit["model_review_receipt"]["author"]["questions"][0]["question_text_spans"] == [
        "Explain the FIRST", " method now."
    ]
    assert audit["original_ids"] == ["QINV-0001", "QINV-0002"]
    assert audit["omitted"] == ["QINV-0002"]
    assert audit["added"] == [added["qid"]]
    assert any(move["identity"] == "QINV-0001" for move in audit["moved"])
    assert audit["model_review_receipt"]["critic"]["verdict"] == "verified"
    assert audit["model_review_receipt"]["author_attempts"][0] == _author_verdict()
    assert audit["model_review_receipt"]["author"] == _author_verdict(resolved=True)
    assert result["round_recorded"] is True
    assert [call["critic"] for call in calls] == [False, True]

    atom = source_inventory.source_atom_from_item(
        items[0],
        source_document_hash=payload["source_document_hash"],
        mined_types=payload["mined_types"],
        type_case_rows=payload["type_case_rows"],
    )
    assert atom["raw_text"] == "Explain the FIRST method now."
    assert atom["normalized_public_text"] == "Explain the FIRST method now."
    assert atom["shared_context"] == "Shared original context."
    assert atom["source_answer"] == "Original answer."
    assert atom["options"] == ["Original option A", "Original option B"]
    assert atom["reviewed_target"]["case_id"] == "CASE-0002"

    example_routes = [
        row for row in payload["type_case_rows"]
        if row.get("row_kind") in {"example", "reviewer_added"}
    ]
    assert {row.get("example_qid") for row in example_routes} == {
        "QINV-0001", added["qid"]
    }
    assert all(row.get("case_id") == "CASE-0002" for row in example_routes)
    assert "QINV-0002" not in json.dumps(payload["type_case_rows"])

    defects = release.structural_defects(payload)
    assert defects == [], defects


def test_unchanged_canonical_three_sheet_roundtrip_preserves_review_bank(
    db, tmp_path: Path, monkeypatch
):
    """An untouched downloaded file must not invoke semantic review or mint a round."""

    job = _canonical_release(db)
    workbook_path = tmp_path / "canonical-unchanged.xlsx"
    workbook_path.write_bytes(
        release_files.build_release_bulk_import_workbook(db, job, lane="post")
    )

    def fail_model(*args, **kwargs):
        raise AssertionError("unchanged canonical workbook triggered model review")

    monkeypatch.setattr(concept_question_review, "_call", fail_model)
    result = release_workbook_edits.apply_workbook_for_review(
        db, job, lane="post", workbook_path=workbook_path
    )
    assert result["round_recorded"] is False
    payload = release.release_payload(job, lane="post")
    assert payload is not None
    from app.services import generation_repair_policy as repair
    assert not repair.active(payload)
    assert [item["qid"] for item in payload["question_task_inventory"]["items"]] == [
        "QINV-0001", "QINV-0002"
    ]
