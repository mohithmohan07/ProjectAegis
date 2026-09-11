"""Step 03 — the reviewed Master file (owner's three-step workflow).

The reviewer downloads a lane's Master, edits it locally, and uploads it
back. The upload is applied VERBATIM as a new release version (contract
v2.0 §38/§44): changed cells are reversed through the exact projection the
renderer wrote them with, omitted rows are recorded omissions, unknown rows
are refused with the Step 02 path named, and the mechanical gates (closed
Q45 vocabulary included) decide readiness — never the upload. Publication
is the separate explicit act: the database write through the one existing
path plus the shared CMS workbook append, idempotent on repeat.
"""
from __future__ import annotations

import copy
import hashlib
import io
from pathlib import Path

import openpyxl
import pytest

from app import config, models
from app.bulk_import import assessment_workbook as aw
from app.services import assessment_release_run as run
from app.services import assessment_release_service as svc
from app.services import build_concepts_release as release
from app.services import build_concepts_release_publication as publication
from app.services import master_review, release_core
from tests.test_assessment_pre_release_lane import SOURCE_INVENTORY
from tests.test_assessment_release_run import (
    OWNER,
    _authorities,
    _chapter_with_concepts,
    _decision_context,
    _make_job,
)
from tests.test_pre_release_lane_wiring import (
    _post_mined_types,
    _routed_post_records,
)


# --------------------------------------------------------------------------- #
# Fixtures: a job whose Post Master was built offline and is ready for review
# --------------------------------------------------------------------------- #

def _routed_post_job(db, chapter) -> models.UploadJob:
    """A Post job whose staged Concept release is mechanically complete.

    The Type/Case routes and the source inventory are the lane-wiring
    fixtures, so the Post Concept file publishes to the database and the
    scripted Master pipeline builds Output 04 from the same two questions.
    """
    job = models.UploadJob(
        owner_sub=OWNER,
        module="build_concepts",
        upload_type="textbook",
        filename="ch.mmd",
        mmd_text="# Chapter\n\nExercise 1. Which of these is a solid?",
        status="generated",
        learning_kind="post",
        source_book="NCERT",
        deposit_scope_type="chapter",
        deposit_scope_ids=[chapter.id],
        question_inventory=copy.deepcopy(SOURCE_INVENTORY),
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    release.stage_release(
        db,
        job,
        target_chapter_id=chapter.id,
        records=_routed_post_records(),
        inventory=copy.deepcopy(SOURCE_INVENTORY),
        mined_types=_post_mined_types(),
        reason="recorded Output-03 fixture",
    )
    db.refresh(job)
    return job


def _master_ready_job(db):
    """One job at the Step 03 boundary: Post Concept staged, Post Master
    published by the offline scripted pipeline, review marker master_ready."""
    chapter = _chapter_with_concepts(db)
    job = _routed_post_job(db, chapter)
    authorities, _ = _authorities(db, chapter)
    published = run.run_release_for_job(
        db, job.id, owner_sub=OWNER, authorities=authorities,
        **_decision_context())
    assert (published.diagnostics or {}).get("readiness") == svc.READY
    release.initialize_concept_review(db, job, target_chapter_id=chapter.id)
    release.accept_concept_review(db, job)
    release.update_concept_review_state(
        db, job,
        status=release.CONCEPT_REVIEW_MASTER_READY,
        master_outputs={
            "post": {"ready": True, "release_id": published.id},
            "pre": {"ready": False, "reason": "no Pre release in this fixture"},
        },
    )
    db.refresh(job)
    return job, published


def _master_bytes(published) -> bytes:
    return svc.published_artifact(published, svc.MASTER_FILENAME)


def _workbook(data: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(data))


def _bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _columns(sheet) -> dict[str, int]:
    return {
        str(cell.value): cell.column
        for cell in sheet[2] if cell.value is not None
    }


def _question_rows(sheet) -> list[int]:
    labels = _columns(sheet)["question_label"]
    return [
        row for row in range(3, sheet.max_row + 1)
        if str(sheet.cell(row=row, column=labels).value or "").strip()
    ]


def _label(sheet, row: int) -> str:
    return str(sheet.cell(row=row, column=_columns(sheet)["question_label"]).value)


def _submit(db, job, data: bytes, *, filename="aegis_master.xlsx", lane="post"):
    return master_review.submit_reviewed_master(
        db, job, lane=lane, workbook_bytes=data, filename=filename,
        owner_sub=OWNER,
    )


def _candidate(release_row, label: str) -> dict:
    return next(
        candidate for candidate in release_row.payload["candidates"]
        if candidate["question_label"] == label
    )


# --------------------------------------------------------------------------- #
# (a) round trip: the unchanged download is not a new version
# --------------------------------------------------------------------------- #

def test_unchanged_master_upload_records_no_round_and_keeps_the_version(db):
    job, published = _master_ready_job(db)
    data = _master_bytes(published)

    result = _submit(db, job, data)

    assert result["round_recorded"] is False
    assert result["version"] == published.version == 1
    assert result["release_id"] == published.id
    assert result["changed_fields"] == []
    assert result["omitted_questions"] == []
    assert result["added_questions"] == []
    assert result["input_sha256"] == hashlib.sha256(data).hexdigest()
    live = release_core.latest_release_for_lane(db, job.id, "post")
    assert live.id == published.id
    db.refresh(job)
    marker = release.concept_review_state(job)
    assert marker["status"] == release.CONCEPT_REVIEW_MASTER_READY
    lane_state = marker["master_review"]["post"]
    assert lane_state["status"] == "reviewed"
    assert lane_state["unchanged"] is True
    assert lane_state["sha256"] == result["input_sha256"]
    assert result["review_workflow"] == marker
    # A re-saved copy with no cell change is the same decision.
    again = _submit(db, job, _bytes(_workbook(data)), filename="resaved.xlsx")
    assert again["round_recorded"] is False
    assert again["version"] == 1


# --------------------------------------------------------------------------- #
# (b) verbatim edits become a new version
# --------------------------------------------------------------------------- #

def test_edited_wording_and_marks_become_a_new_version_recorded_verbatim(db):
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    label = _label(sheet, row)
    before_question = sheet.cell(row=row, column=columns["question"]).value
    assert before_question == "Which of these is a solid?"
    new_wording = "Which of these shapes is a solid?"
    sheet.cell(row=row, column=columns["question"]).value = new_wording
    sheet.cell(row=row, column=columns["marks"]).value = 2
    # The correct option's weight follows the marks (contract §22).
    sheet.cell(row=row, column=columns["answer_weightage_1"]).value = 2
    # The carried column policy keeps the option-label prefix (contract §22).
    sheet.cell(row=row, column=columns["answer_explanation"]).value = (
        "a) Cube. A cube occupies space in all three dimensions.<br>\nIt is a solid."
    )

    result = _submit(db, job, _bytes(workbook), filename="post-master-edited.xlsx")

    assert result["round_recorded"] is True
    assert result["version"] == 2
    assert result["release_uid"] == published.release_uid
    edits = {edit["field"]: edit for edit in result["changed_fields"]}
    assert set(edits) == {"question", "marks", "answer_weightage_1", "answer_explanation"}
    assert edits["question"] == {
        "question_label": label, "field": "question",
        "before": "Which of these is a solid?", "after": new_wording,
    }
    assert edits["marks"]["before"] == 1 and edits["marks"]["after"] == 2
    assert result["omitted_questions"] == []
    assert result["added_questions"] == []
    assert result["readiness"] in {svc.READY, svc.RELEASED_WITH_WARNINGS}, result["issues"]

    new_release = db.get(models.AssessmentRelease, result["release_id"])
    assert new_release.state in {"ready_for_upload", "validated_with_flags"}
    candidate = _candidate(new_release, label)
    # The reviewer's stem is both learner fields (the contract requires
    # question == question_text); the workbook's question_text cell is the
    # re-composition and produced no separate edit.
    assert candidate["question"] == new_wording
    assert candidate["question_text"] == new_wording
    assert candidate["marks"] == 2
    assert candidate["answers"][0]["answer_weightage"] == 2
    # The ``<br>`` projection is inverted to the internal newline.
    assert candidate["answer_explanation"] == (
        "a) Cube. A cube occupies space in all three dimensions.\nIt is a solid."
    )
    # Untouched cells changed nothing on the candidate.
    original = _candidate(published, label)
    for field in ("question_category", "cognitive_skill", "difficulty",
                  "question_duration", "answer_restriction", "source_atom_ids"):
        assert candidate[field] == original[field]
    # The lone blueprint cell follows the reviewed marks (format contract).
    cell = next(
        c for c in new_release.payload["blueprint_cells"]
        if c["cell_id"] == candidate["blueprint_cell_id"]
    )
    assert cell["marks"] == 2
    assert new_release.payload["master_review"]["policy_version"] == (
        master_review.POLICY_VERSION
    )
    assert new_release.payload["master_review"]["sha256"] == result["input_sha256"]
    assert new_release.provider_identity["master_review"]["edit_count"] == 4
    # v2 is published on disk beside v1; v1 is superseded.
    directory = Path(new_release.publication["directory"])
    assert directory.name == "v2"
    assert (directory / svc.MASTER_FILENAME).is_file()
    assert (directory / svc.CONCEPTS_FILENAME).is_file()
    assert (directory / svc.MANIFEST_FILENAME).is_file()
    rendered = aw.parse_workbook((directory / svc.MASTER_FILENAME).read_bytes())
    rendered_row = next(
        r for r in rendered["sheets"]["Objective"]["rows"]
        if r.get("question_label") == label
    )
    assert rendered_row["question"] == new_wording
    assert rendered_row["marks"] == 2
    db.refresh(published)
    assert published.state == "superseded"
    assert release_core.latest_release_for_lane(db, job.id, "post").id == new_release.id

    db.refresh(job)
    marker = release.concept_review_state(job)
    lane_state = marker["master_review"]["post"]
    assert lane_state["status"] == "reviewed"
    assert lane_state["version"] == 2
    assert lane_state["release_id"] == new_release.id
    assert lane_state["filename"] == "post-master-edited.xlsx"
    assert lane_state["changed_fields"] == result["changed_fields"]
    assert lane_state["readiness"] == result["readiness"]
    assert marker["master_outputs"]["post"]["release_id"] == new_release.id
    assert marker["status"] == release.CONCEPT_REVIEW_MASTER_READY

    # Re-uploading the accepted file is not another round.
    repeat = _submit(db, job, _bytes(workbook), filename="post-master-edited.xlsx")
    assert repeat["round_recorded"] is False
    assert repeat["version"] == 2


def test_a_question_text_edit_alone_reduces_to_the_stem(db):
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    label = _label(sheet, row)
    cell = sheet.cell(row=row, column=columns["question_text"])
    assert str(cell.value).startswith("Which of these is a solid?")
    cell.value = str(cell.value).replace(
        "Which of these is a solid?", "Which one of these is a solid?", 1,
    )

    result = _submit(db, job, _bytes(workbook))

    assert result["round_recorded"] is True
    assert [edit["field"] for edit in result["changed_fields"]] == ["question_text"]
    candidate = _candidate(
        db.get(models.AssessmentRelease, result["release_id"]), label)
    assert candidate["question"] == "Which one of these is a solid?"
    assert candidate["question_text"] == candidate["question"]
    assert result["readiness"] in {svc.READY, svc.RELEASED_WITH_WARNINGS}, result["issues"]


# --------------------------------------------------------------------------- #
# (c) a removed row is a recorded omission
# --------------------------------------------------------------------------- #

def test_a_removed_row_is_recorded_as_omitted_and_leaves_the_new_snapshot(db):
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Descriptive"]
    rows = _question_rows(sheet)
    assert len(rows) == 1
    omitted_label = _label(sheet, rows[0])
    omitted_candidate = _candidate(published, omitted_label)
    sheet.delete_rows(rows[0], 1)

    result = _submit(db, job, _bytes(workbook))

    assert result["round_recorded"] is True
    assert result["version"] == 2
    assert result["changed_fields"] == []
    assert [item["question_label"] for item in result["omitted_questions"]] == [
        omitted_label
    ]
    assert result["omitted_questions"][0]["candidate_id"] == (
        omitted_candidate["candidate_id"]
    )
    new_release = db.get(models.AssessmentRelease, result["release_id"])
    labels = {c["question_label"] for c in new_release.concept_snapshot["candidates"]}
    assert omitted_label not in labels
    assert len(labels) == len(published.concept_snapshot["candidates"]) - 1
    # No placement or member list still names the omitted question.
    assert all(
        p["candidate_id"] != omitted_candidate["candidate_id"]
        for p in new_release.payload["placements"]
    )
    assert all(
        omitted_candidate["candidate_id"] not in (g.get("member_candidate_ids") or [])
        for g in new_release.payload["groups"]
    )
    # The emptied variant group is gone (a required shell would stay).
    home = omitted_candidate["group_key"]
    survivors = [g for g in new_release.payload["groups"] if g["group_key"] == home]
    assert all(int(g.get("group_sequence") or 0) == 1 for g in survivors)
    assert new_release.payload["master_review"]["omitted"][0]["question_label"] == omitted_label
    # Rollups are recomputed mechanically by the renderer: the concept's
    # label list and the Descriptive sheet no longer carry the question.
    directory = Path(new_release.publication["directory"])
    rendered = aw.parse_workbook((directory / svc.MASTER_FILENAME).read_bytes())
    assert rendered["sheets"]["Descriptive"]["rows"] == []
    objective_rows = [
        r for r in rendered["sheets"]["Objective"]["rows"] if r.get("question_label")
    ]
    assert objective_rows
    assert all(
        omitted_label not in str(r.get("concept_question_labels") or "")
        for r in objective_rows
    )
    assert result["readiness"] in {svc.READY, svc.RELEASED_WITH_WARNINGS}, result["issues"]


# --------------------------------------------------------------------------- #
# (d) additions are refused with the Step 02 path named
# --------------------------------------------------------------------------- #

def test_an_unknown_label_is_refused_and_points_to_step_two(db):
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    copied = [sheet.cell(row=row, column=c).value for c in range(1, sheet.max_column + 1)]
    copied[columns["question_label"] - 1] = "REVIEWER Q99"
    copied[columns["question"] - 1] = "A reviewer-authored question?"
    sheet.append(copied)

    with pytest.raises(master_review.MasterReviewRefused) as refused:
        _submit(db, job, _bytes(workbook))
    message = str(refused.value)
    assert "'REVIEWER Q99'" in message
    assert "Step 02" in message
    # Nothing was versioned.
    assert release_core.latest_release_for_lane(db, job.id, "post").id == published.id
    assert db.query(models.AssessmentRelease).filter(
        models.AssessmentRelease.release_uid == published.release_uid,
    ).count() == 1


def test_a_blank_label_row_is_refused_the_same_way(db):
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    sheet.cell(row=row, column=columns["question_label"]).value = None
    with pytest.raises(master_review.MasterReviewRefused, match="matches no question"):
        _submit(db, job, _bytes(workbook))


def test_a_moved_sheet_and_a_wrong_layout_are_refused(db):
    job, published = _master_ready_job(db)
    # A question cannot change sheet: the sheet decides the answer layout.
    workbook = _workbook(_master_bytes(published))
    objective = workbook["Objective"]
    subjective = workbook["Subjective"]
    columns = _columns(objective)
    row = _question_rows(objective)[0]
    label = _label(objective, row)
    moved = [None] * subjective.max_column
    sub_columns = _columns(subjective)
    moved[sub_columns["question_label"] - 1] = label
    moved[sub_columns["question"] - 1] = objective.cell(row=row, column=columns["question"]).value
    subjective.append(moved)
    objective.delete_rows(row, 1)
    with pytest.raises(master_review.MasterReviewRefused, match="cannot change sheet"):
        _submit(db, job, _bytes(workbook))

    # A workbook that is not the downloaded layout.
    other = openpyxl.Workbook()
    other.active.title = "Objective"
    other.active.append(["Question"])
    other.active.append(["question_label", "question"])
    with pytest.raises(master_review.MasterReviewRefused, match="downloaded Master layout"):
        _submit(db, job, _bytes(other))

    # Not a workbook at all.
    with pytest.raises(master_review.MasterReviewRefused, match="not a readable"):
        _submit(db, job, b"not a workbook")


# --------------------------------------------------------------------------- #
# (e) the closed CMS vocabulary blocks the release, never the upload
# --------------------------------------------------------------------------- #

def test_a_category_outside_the_closed_vocabulary_blocks_readiness_not_the_upload(
    db, tmp_path, monkeypatch,
):
    monkeypatch.setattr(config, "BULK_IMPORT_OUTPUT", tmp_path / "bulk_import_output.xlsx")
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    label = _label(sheet, row)
    sheet.cell(row=row, column=columns["question_category"]).value = "MCQ"

    result = _submit(db, job, _bytes(workbook))

    assert result["round_recorded"] is True
    assert result["version"] == 2
    assert result["readiness"] == svc.BLOCKED
    assert any(
        "question_category" in issue and "MCQ" in issue for issue in result["issues"]
    ), result["issues"]
    new_release = db.get(models.AssessmentRelease, result["release_id"])
    # The reviewer's value is retained verbatim in the frozen evidence...
    assert _candidate(new_release, label)["question_category"] == "MCQ"
    # ...and the downloadable workbook keeps the cell blank and visibly blocked.
    directory = Path(new_release.publication["directory"])
    rendered = aw.parse_workbook((directory / svc.MASTER_FILENAME).read_bytes())
    rendered_row = next(
        r for r in rendered["sheets"]["Objective"]["rows"]
        if r.get("question_label") == label
    )
    assert rendered_row["question_category"] == ""
    db.refresh(job)
    assert release.concept_review_state(job)["master_review"]["post"]["readiness"] == svc.BLOCKED

    # The publish act is what refuses, naming the block.
    publication.upload_release_to_database(db, job.id, owner_sub=OWNER, lane="post")
    with pytest.raises(master_review.MasterReviewConflict, match="blocked for database upload"):
        master_review.publish_reviewed_master(db, job, lane="post", owner_sub=OWNER)
    assert db.query(models.Question).filter(
        models.Question.question_label == label,
    ).count() == 0


# --------------------------------------------------------------------------- #
# (f) publish: database rows AND the shared CMS workbook, idempotently
# --------------------------------------------------------------------------- #

def _workbook_labels(path: Path) -> dict[str, list[str]]:
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        found: dict[str, list[str]] = {}
        for sheet in workbook.sheetnames:
            ws = workbook[sheet]
            header = [
                str(c or "") for c in next(
                    ws.iter_rows(min_row=2, max_row=2, values_only=True), (),
                )
            ]
            if "question_label" not in header:
                continue
            index = header.index("question_label")
            found[sheet] = [
                str(row[index])
                for row in ws.iter_rows(min_row=3, values_only=True)
                if row and index < len(row) and str(row[index] or "").strip()
            ]
        return found
    finally:
        workbook.close()


def test_publish_requires_the_concept_lane_then_writes_database_and_cms_workbook(
    db, tmp_path, monkeypatch,
):
    target = tmp_path / "bulk_import_output.xlsx"
    monkeypatch.setattr(config, "BULK_IMPORT_OUTPUT", target)
    job, published = _master_ready_job(db)
    labels = [c["question_label"] for c in published.concept_snapshot["candidates"]]
    assert len(labels) == 2

    with pytest.raises(master_review.MasterReviewConflict, match="Publish the post Concept file first"):
        master_review.publish_reviewed_master(db, job, lane="post", owner_sub=OWNER)
    assert db.query(models.Question).filter(
        models.Question.question_label.in_(labels)).count() == 0

    concept = publication.upload_release_to_database(
        db, job.id, owner_sub=OWNER, lane="post")
    assert concept["database_uploaded"] is True
    assert target.is_file()
    before = _workbook_labels(target)
    assert not any(label in rows for rows in before.values() for label in labels)

    result = master_review.publish_reviewed_master(db, job, lane="post", owner_sub=OWNER)

    assert result["lane"] == "post"
    assert result["release_id"] == published.id
    assert result["version"] == 1
    assert result["publication_status"] == "published"
    assert result["database"]["questions_created"] == 2
    assert result["database"]["release_uid"] == published.release_uid
    assert result["cms_workbook"]["status"] == "published"
    assert result["cms_workbook"]["path"] == target.name
    assert result["cms_workbook"]["objective"] == 1
    assert result["cms_workbook"]["descriptive"] == 1
    rows = db.query(models.Question).filter(
        models.Question.question_label.in_(labels)).all()
    assert len(rows) == 2
    assert all(row.route_audit["release_uid"] == published.release_uid for row in rows)
    assert all(row.group.concept_id for row in rows)
    after = _workbook_labels(target)
    assert set(labels) <= {label for rows in after.values() for label in rows}
    assert not Path(str(target) + ".outbox.json").exists()
    db.refresh(published)
    assert published.state == "uploaded"

    db.refresh(job)
    marker = release.concept_review_state(job)
    lane_state = marker["master_review"]["post"]
    assert lane_state["status"] == "published"
    assert lane_state["published"]["database"]["questions_created"] == 2
    assert lane_state["published"]["cms_workbook"]["path"] == target.name
    assert lane_state["published"]["version"] == 1
    assert marker["status"] == release.CONCEPT_REVIEW_PUBLISHED
    assert result["review_workflow"] == marker

    # Idempotent: the recorded receipt comes back, nothing is written twice.
    again = master_review.publish_reviewed_master(db, job, lane="post", owner_sub=OWNER)
    assert again["database"] == result["database"]
    assert again["cms_workbook"] == result["cms_workbook"]
    assert again["publication_status"] == "published"
    assert db.query(models.Question).filter(
        models.Question.question_label.in_(labels)).count() == 2
    assert _workbook_labels(target) == after


def test_publish_takes_the_reviewed_version_when_one_was_accepted(
    db, tmp_path, monkeypatch,
):
    target = tmp_path / "bulk_import_output.xlsx"
    monkeypatch.setattr(config, "BULK_IMPORT_OUTPUT", target)
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    label = _label(sheet, row)
    sheet.cell(row=row, column=columns["question"]).value = "Which of these objects is a solid?"
    reviewed = _submit(db, job, _bytes(workbook))
    assert reviewed["version"] == 2

    publication.upload_release_to_database(db, job.id, owner_sub=OWNER, lane="post")
    result = master_review.publish_reviewed_master(db, job, lane="post", owner_sub=OWNER)

    assert result["version"] == 2
    assert result["release_id"] == reviewed["release_id"]
    stored = db.query(models.Question).filter(
        models.Question.question_label == label).one()
    assert stored.question == "Which of these objects is a solid?"
    assert stored.route_audit["version"] == 2
    assert label in {l for rows in _workbook_labels(target).values() for l in rows}


# --------------------------------------------------------------------------- #
# (g) the routes: 400 / 404 / 409 / 422 and the response shapes
# --------------------------------------------------------------------------- #

def _post_file(client, url: str, data: bytes, name: str = "aegis_master.xlsx"):
    return client.post(
        url,
        files={"file": (
            name, io.BytesIO(data),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
    )


def test_routes_refuse_missing_jobs_legacy_jobs_and_early_states(client, db):
    response = _post_file(
        client, "/build-concepts/uploads/999999/master-review/submit?lane=post", b"x")
    assert response.status_code == 404
    assert client.post(
        "/build-concepts/uploads/999999/master-review/publish?lane=post"
    ).status_code == 404

    # The lane is never defaulted on a publication-shaped act.
    chapter = _chapter_with_concepts(db)
    legacy = _make_job(db, chapter)
    response = _post_file(
        client, f"/build-concepts/uploads/{legacy.id}/master-review/submit", b"x")
    assert response.status_code == 400
    assert "lane=post" in response.json()["detail"]

    # A legacy job without the Concept review marker keeps its old routes.
    response = _post_file(
        client, f"/build-concepts/uploads/{legacy.id}/master-review/submit?lane=post", b"x")
    assert response.status_code == 409
    assert "no Concept review gate" in response.json()["detail"]
    response = client.post(
        f"/build-concepts/uploads/{legacy.id}/master-review/publish?lane=post")
    assert response.status_code == 409
    assert "no Concept review gate" in response.json()["detail"]

    # A job still under Concept review is not at Step 03.
    release.initialize_concept_review(db, legacy, target_chapter_id=chapter.id)
    response = _post_file(
        client, f"/build-concepts/uploads/{legacy.id}/master-review/submit?lane=post", b"x")
    assert response.status_code == 409
    assert "pending_review" in response.json()["detail"]
    response = client.post(
        f"/build-concepts/uploads/{legacy.id}/master-review/publish?lane=post")
    assert response.status_code == 409


def test_routes_apply_refuse_and_publish_end_to_end(client, db, tmp_path, monkeypatch):
    target = tmp_path / "bulk_import_output.xlsx"
    monkeypatch.setattr(config, "BULK_IMPORT_OUTPUT", target)
    job, published = _master_ready_job(db)
    base = f"/build-concepts/uploads/{job.id}/master-review"

    # A lane with no Master release is 404.
    response = _post_file(client, f"{base}/submit?lane=pre", _master_bytes(published))
    assert response.status_code == 404
    assert "no live pre Master release" in response.json()["detail"]

    # Wrong layout is a readable 422.
    other = openpyxl.Workbook()
    other.active.append(["nothing"])
    response = _post_file(client, f"{base}/submit?lane=post", _bytes(other))
    assert response.status_code == 422
    assert "downloaded Master layout" in response.json()["detail"]

    # Unknown label is a readable 422 naming Step 02.
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    copied = [sheet.cell(row=row, column=c).value for c in range(1, sheet.max_column + 1)]
    copied[columns["question_label"] - 1] = "ROUTE Q77"
    sheet.append(copied)
    response = _post_file(client, f"{base}/submit?lane=post", _bytes(workbook))
    assert response.status_code == 422
    assert "Step 02" in response.json()["detail"]

    # A verbatim edit through the route.
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    sheet.cell(row=row, column=columns["level_of_difficulty"]).value = "Moderate"
    response = _post_file(client, f"{base}/submit?lane=post", _bytes(workbook), "edited.xlsx")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["round_recorded"] is True
    assert body["version"] == 2
    assert body["filename"] == "edited.xlsx"
    assert body["lane"] == "post"
    assert [edit["field"] for edit in body["changed_fields"]] == ["level_of_difficulty"]
    assert body["changed_fields"][0]["before"] == "Less"
    assert body["changed_fields"][0]["after"] == "Moderate"
    assert body["master_review"]["version"] == 2
    assert body["review_workflow"]["master_review"]["post"]["version"] == 2
    assert set(body) >= {
        "lane", "filename", "input_sha256", "release_id", "release_uid", "version",
        "round_recorded", "changed_fields", "omitted_questions", "added_questions",
        "readiness", "issues", "master_review", "review_workflow",
    }

    # Publish: 409 before the Concept lane, then the full receipt.
    response = client.post(f"{base}/publish?lane=post")
    assert response.status_code == 409
    assert "Publish the post Concept file first" in response.json()["detail"]
    assert client.post(
        f"/build-concepts/uploads/{job.id}/upload-release?lane=post"
    ).status_code == 200
    response = client.post(f"{base}/publish?lane=post")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["publication_status"] == "published"
    assert body["version"] == 2
    assert body["database"]["questions_created"] == 2
    assert body["cms_workbook"]["status"] == "published"
    assert body["master_review"]["status"] == "published"
    assert body["review_workflow"]["status"] == release.CONCEPT_REVIEW_PUBLISHED
    assert set(body) >= {
        "lane", "release_id", "release_uid", "version", "database", "cms_workbook",
        "publication_status", "master_review", "review_workflow",
    }
    # Repeating through the route is the recorded receipt.
    again = client.post(f"{base}/publish?lane=post")
    assert again.status_code == 200
    assert again.json()["database"] == body["database"]


# --------------------------------------------------------------------------- #
# Marker mechanics
# --------------------------------------------------------------------------- #

def test_marker_defaults_and_status_vocabulary():
    assert release.CONCEPT_REVIEW_PUBLISHED == "published"
    assert release.CONCEPT_REVIEW_PUBLISHED in release.CONCEPT_REVIEW_STATUSES
    job = models.UploadJob(
        module="build_concepts",
        question_inventory={release.CONCEPT_REVIEW_KEY: {"status": "master_ready"}},
    )
    assert release.concept_review_state(job)["master_review"] == {}
    assert release.concept_review_state(
        models.UploadJob(module="build_concepts", question_inventory={"items": []})
    ) == {}


def test_master_review_state_merges_per_lane(db):
    job, _published = _master_ready_job(db)
    first = release.update_concept_review_state(
        db, job, master_review={"post": {"status": "reviewed", "version": 2}},
    )
    assert first["master_review"]["post"] == {"status": "reviewed", "version": 2}
    second = release.update_concept_review_state(
        db, job, master_review={"pre": {"status": "reviewed", "version": 1}},
    )
    assert second["master_review"] == {
        "post": {"status": "reviewed", "version": 2},
        "pre": {"status": "reviewed", "version": 1},
    }
    assert second["status"] == release.CONCEPT_REVIEW_MASTER_READY


# --------------------------------------------------------------------------- #
# Reverse-projection details across rounds and sheets
# --------------------------------------------------------------------------- #

def test_a_second_round_keeps_the_first_rounds_wording_despite_a_stale_question_text(db):
    """Excel users edit ``question`` and leave the derived ``question_text``.

    Round 2 re-uploads round 1's file with one more edit: the stale
    composition must be recognised as an earlier version's cell, never
    reversed into a silent revert of the accepted wording.
    """
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    label = _label(sheet, row)
    sheet.cell(row=row, column=columns["question"]).value = "Which of these shapes is a solid?"
    first = _submit(db, job, _bytes(workbook), filename="round-1.xlsx")
    assert first["version"] == 2

    # Round 2: the same local file, one more cell edited.
    sheet.cell(row=row, column=columns["question_duration"]).value = 3
    second = _submit(db, job, _bytes(workbook), filename="round-2.xlsx")

    assert second["round_recorded"] is True
    assert second["version"] == 3
    assert [edit["field"] for edit in second["changed_fields"]] == ["question_duration"]
    assert any("earlier version" in flag for flag in second["issues"])
    candidate = _candidate(db.get(models.AssessmentRelease, second["release_id"]), label)
    assert candidate["question"] == "Which of these shapes is a solid?"
    assert candidate["question_text"] == candidate["question"]
    assert candidate["question_duration"] == 3
    assert second["readiness"] in {svc.READY, svc.RELEASED_WITH_WARNINGS}, second["issues"]
    db.refresh(job)
    lane_state = release.concept_review_state(job)["master_review"]["post"]
    assert lane_state["rounds"] == 2
    assert lane_state["version"] == 3


def test_descriptive_rubric_and_model_answer_edits_reverse_into_the_candidate(db):
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Descriptive"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    label = _label(sheet, row)
    original = _candidate(published, label)
    assert len(original["answers"]) == 3
    assert sheet.cell(row=row, column=columns["answer_type_1"]).value == "Phrases"
    sheet.cell(row=row, column=columns["answer_content_1"]).value = (
        "all three dimensions are named explicitly"
    )
    sheet.cell(row=row, column=columns["display_answer"]).value = (
        "A cube occupies space in all three dimensions.<br>\nIt has length, breadth and height."
    )
    sheet.cell(row=row, column=columns["answer_explanation"]).value = (
        "A cube occupies space in all three dimensions.<br>\nIt has length, breadth and height."
    )
    # Typing the contract literal in a different accepted spelling is not an edit.
    sheet.cell(row=row, column=columns["answer_type_2"]).value = "phrases"

    result = _submit(db, job, _bytes(workbook))

    assert result["round_recorded"] is True
    fields = sorted(edit["field"] for edit in result["changed_fields"])
    assert fields == ["answer_content_1", "answer_explanation", "display_answer"]
    candidate = _candidate(db.get(models.AssessmentRelease, result["release_id"]), label)
    assert candidate["answers"][0]["answer_content"] == "all three dimensions are named explicitly"
    assert candidate["answers"][0]["answer_weightage"] == original["answers"][0]["answer_weightage"]
    assert candidate["answers"][1:] == original["answers"][1:]
    assert candidate["display_answer"] == (
        "A cube occupies space in all three dimensions.\nIt has length, breadth and height."
    )
    assert candidate["answer_explanation"] == candidate["display_answer"]
    assert candidate["question"] == original["question"]
    assert result["readiness"] in {svc.READY, svc.RELEASED_WITH_WARNINGS}, result["issues"]


def test_blanking_a_rubric_block_removes_that_criterion(db):
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Descriptive"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    label = _label(sheet, row)
    original = _candidate(published, label)
    for cell in ("answer_type_3", "answer_weightage_3", "answer_content_3"):
        sheet.cell(row=row, column=columns[cell]).value = None
    # The category fixes the marks at 3, so the two surviving criteria share
    # them in half-mark steps (contract §24 / Q33).
    sheet.cell(row=row, column=columns["answer_weightage_1"]).value = 1.5
    sheet.cell(row=row, column=columns["answer_weightage_2"]).value = 1.5

    result = _submit(db, job, _bytes(workbook))

    assert result["round_recorded"] is True
    assert {edit["field"] for edit in result["changed_fields"]} == {
        "answer_type_3", "answer_weightage_3", "answer_content_3",
        "answer_weightage_1", "answer_weightage_2",
    }
    candidate = _candidate(db.get(models.AssessmentRelease, result["release_id"]), label)
    assert len(candidate["answers"]) == 2
    assert [a["answer_content"] for a in candidate["answers"]] == [
        a["answer_content"] for a in original["answers"][:2]
    ]
    assert [a["answer_weightage"] for a in candidate["answers"]] == [1.5, 1.5]
    assert candidate["marks"] == original["marks"]
    removed = next(
        edit for edit in result["changed_fields"] if edit["field"] == "answer_content_3"
    )
    assert removed["before"] == original["answers"][2]["answer_content"]
    assert removed["after"] == ""
    assert result["readiness"] in {svc.READY, svc.RELEASED_WITH_WARNINGS}, result["issues"]


def test_group_description_edits_reach_the_group_and_identity_cells_do_not(db):
    job, published = _master_ready_job(db)
    workbook = _workbook(_master_bytes(published))
    sheet = workbook["Objective"]
    columns = _columns(sheet)
    row = _question_rows(sheet)[0]
    label = _label(sheet, row)
    group_key = _candidate(published, label)["group_key"]
    sheet.cell(row=row, column=columns["group_description"]).value = (
        "Identifying a solid among plane and solid shapes."
    )
    sheet.cell(row=row, column=columns["group_name"]).value = "Renamed by hand"
    sheet.cell(row=row, column=columns["concept_details"]).value = "Edited concept prose"

    result = _submit(db, job, _bytes(workbook))

    assert result["round_recorded"] is True
    assert result["changed_fields"] == []
    new_release = db.get(models.AssessmentRelease, result["release_id"])
    group = next(g for g in new_release.payload["groups"] if g["group_key"] == group_key)
    assert group["semantic_description"] == "Identifying a solid among plane and solid shapes."
    # The visible names are the group id (SOP §6.1) and the concept band is
    # the Concept file's: both changes are recorded as not applied.
    assert group["group_name"] == group_key
    assert any("group_name" in flag and "not applied" in flag for flag in result["issues"])
    assert any("concept_details" in flag and "Step 02" in flag for flag in result["issues"])
    assert new_release.payload["master_review"]["group_edits"][0]["group_key"] == group_key
    assert result["readiness"] in {svc.READY, svc.RELEASED_WITH_WARNINGS}, result["issues"]
