"""Focused Q41 Concept review handoff tests."""
from __future__ import annotations

import asyncio
import copy
import io
import sys
import threading
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services import build_concepts_release as release
from app.services import build_concepts_release_contract as release_contract
from app.services import build_concepts_release_api_contract as release_api
from app.services import build_concepts_release_files as release_files
from app.services import auth, openai_usage, run_state, uploads
from app.services import release_workbook_edits, reviewed_file_input
from app.services.concept_question_review import ReviewRows

sys.path.insert(0, str(Path(__file__).parent))
from test_build_concepts_release import (
    _inventory,
    _job,
    _mined_types,
    _rendered_records,
)


@pytest.fixture(autouse=True)
def _reset_progress_floor():
    yield
    # Direct lifecycle calls emit progress in the caller context. Keep this
    # module's route/lifecycle tests from affecting another module's progress
    # allocation assertions.
    from app.services import progress
    progress._progress_floor.set(0.0)


def test_corrected_routing_workbook_replaces_exact_reviewed_post_set(db, tmp_path: Path):
    job, chapter = _job(db)
    release.stage_release(
        db,
        job,
        target_chapter_id=chapter.id,
        records=_rendered_records(),
        inventory=_inventory(),
        mined_types=_mined_types(),
    )
    workbook = load_workbook(
        io.BytesIO(release_files.build_release_workbook(job))
    )
    routes = workbook["Type Case Routing"]
    # Edit visible Example wording, move its owner topic, and omit Q2 by
    # deleting that Example row.  Add a deliberate reviewer row as well.
    routes.cell(4, 10).value = "Reviewer wording for Q1"
    routes.cell(4, 7).value = "TOPIC-B"
    routes.delete_rows(6, 1)
    routes.append([
        "reviewer_added", "", "", "", "CASE-NEW", "", "TOPIC-A", "",
        "", "Reviewer added question", "ready", "",
    ])
    path = tmp_path / "corrected.xlsx"
    workbook.save(path)

    result = release_workbook_edits.apply_workbook_for_review(
        db, job, lane="post", workbook_path=path, owner_sub="test-owner"
    )
    db.refresh(job)
    payload = release.release_payload(job, lane="post")
    assert payload is not None
    items = payload["question_task_inventory"]["items"]
    assert [row["qid"] for row in items] == ["QINV-0001", next(
        row["qid"] for row in items if row.get("provenance") == "reviewer_added"
    )]
    assert items[0]["raw_task"] == "Reviewer wording for Q1"
    assert items[0]["topic_hint"] == "TOPIC-B"
    audit = payload["question_task_inventory"]["reviewed_source_questions"]
    assert audit["omitted"] == ["QINV-0002"]
    assert len(audit["added"]) == 1
    assert result["round_recorded"] is True


def test_empty_routing_surface_is_recorded_as_remove_all(db, tmp_path: Path):
    job, chapter = _job(db)
    release.stage_release(
        db,
        job,
        target_chapter_id=chapter.id,
        records=_rendered_records(),
        inventory=_inventory(),
        mined_types=_mined_types(),
    )
    workbook = load_workbook(
        io.BytesIO(release_files.build_release_workbook(job))
    )
    routes = workbook["Type Case Routing"]
    for row in range(routes.max_row, 1, -1):
        if routes.cell(row, 1).value == "example":
            routes.delete_rows(row, 1)
    path = tmp_path / "remove-all.xlsx"
    workbook.save(path)
    release_workbook_edits.apply_workbook_for_review(
        db, job, lane="post", workbook_path=path
    )
    payload = release.release_payload(job, lane="post")
    assert payload is not None
    assert payload["question_task_inventory"]["items"] == []
    assert payload["question_task_inventory"]["reviewed_source_questions"]["omitted"] == [
        "QINV-0001", "QINV-0002"
    ]


def test_explicit_question_replacement_clears_stale_source_dependencies():
    payload = {
        "question_task_inventory": {
            "items": [{
                "qid": "Q1",
                "raw_task": "Old task",
                "normalized_task": "Old task",
                "polished_task": "Old task",
                "options": ["old option"],
                "shared_context": "Old context",
                "raw_solution_or_answer": "Old answer",
                "image_urls": ["https://example.test/old.png"],
                "image_manifest": [{"url": "https://example.test/old.png"}],
                "tables": [{"cells": [["old"]]}],
                "content_objects": {"table": "old"},
                "compound_subparts": [{"source_qid": "Q1.1"}],
                "source_context": {"old": "evidence"},
            }],
        },
        "type_case_rows": [],
    }
    reviewed = ReviewRows([{
        "row": "Concept Details:3",
        "kind": "source",
        "question_id": "Q1",
        "source_qid": "Q1",
        "question_text": "New task with replacement evidence",
        "topic_hint": "Topic A",
        "options": [],
        "shared_context": "New context",
        "raw_solution_or_answer": "New answer",
        "preserve_source_dependencies": False,
        "type_id": "",
        "type_title": "",
        "type_definition": "",
        "case_id": "",
        "case_definition": "",
    }], receipt={})

    release_workbook_edits._apply_question_review(
        payload, reviewed, lane="post"
    )
    item = payload["question_task_inventory"]["items"][0]
    assert item["raw_task"] == "New task with replacement evidence"
    assert item["source_context"] == {
        "source_answer": "New answer",
        "shared_context": "New context",
        "options": [],
    }
    for field in (
        "image_urls", "image_manifest", "tables", "content_objects",
        "compound_subparts",
    ):
        assert field not in item
    assert payload["review_question_audit"]["model_review_receipt"] == {}


def test_pre_review_regeneration_uses_corrected_pre_release(tmp_path, db, monkeypatch):
    job, chapter = _job(db)
    release.stage_release(
        db,
        job,
        target_chapter_id=chapter.id,
        records=_rendered_records(),
        inventory=_inventory(),
        mined_types=_mined_types(),
    )
    release.stage_pre_release(
        db,
        job,
        target_chapter_id=chapter.id,
        pre_map={"rows": [{"_pre_concept_id": "PRE-1", "concept_details": "Old"}]},
        pre_questions={"questions": {"PRE-1": [{"pre_question_id": "OLD"}]}},
        inventory={"items": []},
    )
    release.initialize_concept_review(db, job, target_chapter_id=chapter.id)
    db.refresh(job)
    state = release.update_concept_review_state(
        db,
        job,
        reviewed_lane="pre",
        corrected_filename="edited-pre.xlsx",
        corrected_changed=True,
    )
    assert state["corrected_inputs"]["pre"]["filename"] == "edited-pre.xlsx"
    pre_payload = release.release_payload(job, lane=release.LANE_PRE)
    assert pre_payload is not None
    edited = copy.deepcopy(pre_payload["records"])
    edited[0]["concept_details"] = "Corrected prerequisite base"
    durable = copy.deepcopy(dict(job.question_inventory or {}))
    pre_payload["records"] = edited
    durable[release.PRE_RELEASE_KEY] = pre_payload
    job.question_inventory = durable
    db.commit()
    db.refresh(job)

    artifact_dir = tmp_path / "artifacts"
    artifact_dir.mkdir()
    (artifact_dir / "source.phase3-envelope.json").write_text(
        '{"envelope": {"envelope_sha256": "' + ("e" * 64) + '"}}',
        encoding="utf-8",
    )
    monkeypatch.setattr(
        release_contract.uploads,
        "source_artifact_directory",
        lambda _job_id: artifact_dir,
    )
    # The envelope validator is imported inside the function, so patch the
    # module object at its actual import location.
    from app.services.phase3 import envelope as phase3_envelope
    from app.services.phase3 import prequestions
    monkeypatch.setattr(phase3_envelope, "validate", lambda value: value)
    seen: dict[str, object] = {}

    def fake_build(env, pre_map, *, store):
        seen["env"] = env
        seen["rows"] = pre_map["rows"]
        return {"questions": {"PRE-1": [{"pre_question_id": "NEW"}]}}

    monkeypatch.setattr(prequestions, "build", fake_build)
    result = release_contract._regenerate_pre_questions_after_review(db, job)

    assert result["pre_questions_count"] == 1
    assert seen["rows"][0]["concept_details"] == "Corrected prerequisite base"
    db.refresh(job)
    accepted = release.release_payload(job, lane=release.LANE_PRE)
    assert accepted["records"][0]["concept_details"] == "Corrected prerequisite base"
    assert [question["pre_question_id"] for question in accepted["generated_questions"]] == ["NEW"]
    assert accepted[release.STAGED_RELEASE_UID_FIELD] == result["pre_questions_regenerated_for_uid"]
    monkeypatch.setattr(prequestions, "build", lambda *args, **kwargs: pytest.fail("accepted bank regenerated on retry"))
    assert release_contract._regenerate_pre_questions_after_review(db, job) is None


def _review_job(db):
    job, _chapter = _job(db)
    job.status = "concept_review"
    job.question_inventory = {
        "items": [],
        release.CONCEPT_REVIEW_KEY: {
            "version": release.CONCEPT_REVIEW_VERSION,
            "status": release.CONCEPT_REVIEW_PENDING,
            "available_lanes": [release.LANE_POST],
            "required_lanes": [release.LANE_POST],
            "optional_lanes": [],
            "reviewed_lanes": [],
            "concept_versions": {},
            "concept_release_uids": {},
            "corrected_inputs": {},
        },
    }
    db.commit()
    db.refresh(job)
    uploads.start_or_resume_run(
        db,
        job.id,
        owner_sub=auth.LOCAL_OWNER_SUB,
        stage="Concept files ready for review",
        progress_value=0.70,
    )
    uploads.pause_run_for_review(
        db,
        job.id,
        owner_sub=auth.LOCAL_OWNER_SUB,
        progress_value=0.70,
    )
    db.refresh(job)
    return job


def _usage_response():
    usage = SimpleNamespace(
        prompt_tokens=100,
        completion_tokens=20,
        total_tokens=120,
        prompt_tokens_details=SimpleNamespace(cached_tokens=0),
        completion_tokens_details=SimpleNamespace(reasoning_tokens=0),
    )
    return SimpleNamespace(
        model="gpt-5.4-mini-2026-03-17",
        usage=usage,
        choices=[SimpleNamespace(
            message=SimpleNamespace(content="{}"), finish_reason="stop"
        )],
    )


def test_corrected_upload_model_round_uses_same_run_and_pauses_again(
    db, monkeypatch
):
    job = _review_job(db)
    caller_thread = threading.get_ident()
    seen: dict[str, object] = {}

    async def read_upload(_file, *, description):
        assert description == "corrected Concept workbook"
        return b"placeholder workbook"

    def apply_review(worker_db, worker_job, **kwargs):
        seen["worker_thread"] = threading.get_ident()
        seen["run_id"] = worker_job.run_id
        openai_usage.record_response(_usage_response())
        return {"round_recorded": True}

    monkeypatch.setattr(release_api, "read_limited_upload", read_upload)
    monkeypatch.setattr(
        reviewed_file_input,
        "queue",
        apply_review,
    )
    from fastapi import UploadFile

    result = asyncio.run(release_api._concept_review_upload_endpoint(
        job.id,
        lane="post",
        file=UploadFile(io.BytesIO(b"input"), filename="edited.xlsx"),
        db=db,
        user=auth.LOCAL_PRINCIPAL,
    ))

    assert seen["worker_thread"] != caller_thread
    assert seen["run_id"] == job.run_id
    assert result["openai_usage"]["request_count"] == 1
    assert result["concept_review"]["status"] == release.CONCEPT_REVIEW_REVIEWED
    assert result["run_state"]["status"] == "review"
    db.expire_all()
    saved = uploads.get_job(db, job.id)
    assert saved.openai_usage["request_count"] == 1
    assert any(
        event.get("type") == "log" and "API estimate" in event.get("message", "")
        for event in saved.generation_log
    )
    state = run_state.for_job(saved)
    assert state["status"] == "review"
    assert state["active_started_at_epoch"] is None
    assert state["review_started_at_epoch"] is not None
    assert state["progress"] == 0.70


def test_corrected_upload_failure_persists_usage_and_run_history(
    db, monkeypatch
):
    job = _review_job(db)

    async def read_upload(_file, *, description):
        return b"placeholder workbook"

    def failed_review(*_args, **_kwargs):
        openai_usage.record_response(_usage_response())
        raise RuntimeError("review provider failed")

    monkeypatch.setattr(release_api, "read_limited_upload", read_upload)
    monkeypatch.setattr(
        reviewed_file_input,
        "queue",
        failed_review,
    )
    from fastapi import HTTPException, UploadFile

    with pytest.raises(HTTPException) as caught:
        asyncio.run(release_api._concept_review_upload_endpoint(
            job.id,
            lane="post",
            file=UploadFile(io.BytesIO(b"input"), filename="edited.xlsx"),
            db=db,
            user=auth.LOCAL_PRINCIPAL,
        ))

    assert caught.value.status_code == 500
    assert "could not finish the corrected-file upload response" in (
        caught.value.detail
    )
    assert "review provider failed" not in caught.value.detail
    assert isinstance(caught.value.__cause__, RuntimeError)
    assert str(caught.value.__cause__) == "review provider failed"

    db.expire_all()
    saved = uploads.get_job(db, job.id)
    assert saved.openai_usage["request_count"] == 1
    assert saved.openai_usage["total_tokens"] == 120
    assert any(
        event.get("error", {}).get("exception_type") == "RuntimeError"
        for event in saved.generation_log
    )
    state = run_state.for_job(saved)
    assert state["status"] == "review"
    assert state["review_started_at_epoch"] is not None
    assert state["active_started_at_epoch"] is None


def test_corrected_upload_retry_keeps_run_id_and_adds_only_new_receipt(
    db, monkeypatch
):
    job = _review_job(db)
    attempts = 0

    async def read_upload(_file, *, description):
        return b"placeholder workbook"

    def apply_review(worker_db, worker_job, **kwargs):
        nonlocal attempts
        attempts += 1
        openai_usage.record_response(_usage_response())
        if attempts == 1:
            raise RuntimeError("transient review failure")
        return {"round_recorded": True}

    monkeypatch.setattr(release_api, "read_limited_upload", read_upload)
    monkeypatch.setattr(
        reviewed_file_input,
        "queue",
        apply_review,
    )
    from fastapi import HTTPException, UploadFile

    with pytest.raises(HTTPException) as caught:
        asyncio.run(release_api._concept_review_upload_endpoint(
            job.id,
            lane="post",
            file=UploadFile(io.BytesIO(b"input"), filename="edited.xlsx"),
            db=db,
            user=auth.LOCAL_PRINCIPAL,
        ))
    assert caught.value.status_code == 500
    assert "could not finish the corrected-file upload response" in (
        caught.value.detail
    )
    assert "transient review failure" not in caught.value.detail
    assert isinstance(caught.value.__cause__, RuntimeError)
    assert str(caught.value.__cause__) == "transient review failure"
    db.expire_all()
    first = uploads.get_job(db, job.id)
    run_id = first.run_id
    assert first.openai_usage["request_count"] == 1

    result = asyncio.run(release_api._concept_review_upload_endpoint(
        job.id,
        lane="post",
        file=UploadFile(io.BytesIO(b"input"), filename="edited-again.xlsx"),
        db=db,
        user=auth.LOCAL_PRINCIPAL,
    ))
    assert result["openai_usage"]["request_count"] == 2
    db.expire_all()
    second = uploads.get_job(db, job.id)
    assert second.run_id == run_id
    assert second.openai_usage["request_count"] == 2
    assert run_state.for_job(second)["status"] == "review"


def test_master_building_marker_recovers_after_worker_crash(db, monkeypatch):
    job, chapter = _job(db)
    release.stage_release(
        db,
        job,
        target_chapter_id=chapter.id,
        records=_rendered_records(),
        inventory=_inventory(),
        mined_types=_mined_types(),
    )
    release.stage_pre_release(
        db,
        job,
        target_chapter_id=chapter.id,
        pre_map={"rows": []},
        pre_questions={"questions": {}},
        inventory={"items": []},
    )
    release.initialize_concept_review(db, job, target_chapter_id=chapter.id)
    release.accept_concept_review(db, job)

    def crash(*_args, **_kwargs):
        # KeyboardInterrupt models a process/worker crash: the lifecycle
        # marker is committed before this call and the broad Exception handler
        # must not turn the crash into a terminal, non-retryable state.
        raise KeyboardInterrupt("worker crashed")

    monkeypatch.setattr(
        release_api.release_contract,
        "_build_master_siblings",
        crash,
    )
    with pytest.raises(KeyboardInterrupt, match="worker crashed"):
        release_api.release_contract.build_review_masters(
            db, job.id, owner_sub=auth.LOCAL_OWNER_SUB
        )

    db.expire_all()
    crashed = uploads.get_job(db, job.id)
    assert release.concept_review_state(crashed)["status"] == (
        release.CONCEPT_REVIEW_MASTER_BUILDING
    )

    monkeypatch.setattr(
        release_api.release_contract,
        "_build_master_siblings",
        lambda *_args, **_kwargs: {
            release.LANE_PRE: {"release_id": 101},
            release.LANE_POST: {"release_id": 102},
        },
    )
    result = release_api.release_contract.build_review_masters(
        db, job.id, owner_sub=auth.LOCAL_OWNER_SUB
    )
    assert result["all_four_outputs_ready"] is True
    db.expire_all()
    recovered = uploads.get_job(db, job.id)
    assert release.concept_review_state(recovered)["status"] == (
        release.CONCEPT_REVIEW_MASTER_READY
    )
